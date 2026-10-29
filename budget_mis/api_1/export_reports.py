"""Excel export for the Budget vs Actuals page (frontend/src/pages/
BudgetVsActual.vue) - api_1's replacement for
budget_mis.api.export_reports.export_budget_vs_actual.

This is a pure relocation, not a correctness fix: unlike the rest of this
session's api -> api_1 migration, export_budget_vs_actual itself has no bug
to fix and nothing to cache or scope - it only serializes heads_data, the
already-computed tree the frontend got back from the (now properly
unit-scoped and cached) api_1.budget_vs_actual.get_combined_actuals_report,
exactly as the user already sees it on screen. There is no re-fetch, no
permission check of its own to add: whatever that report call already
scoped the data to is what gets exported.

The shared Excel styling constants (_BOLD, _FILL_BLUE_MID, NUM_FMT_RAW,
etc.) and _fv/_wb_to_b64 helpers are imported directly from
budget_mis.api.export_reports rather than duplicated here - they're reused
by every other export sheet in that module (Annual, Estimate, PPT,
Headcount, Summary INR...), so copying them would just create a second
copy that can silently drift from the original's own styling choices.
"""

import json

from openpyxl import Workbook
from openpyxl.styles import Font

from budget_mis.utils import guest_api
from budget_mis.api.export_reports import (
	_fv,
	_wb_to_b64,
	_BOLD,
	_WHITE_BOLD,
	_FILL_BLUE_MID,
	_FILL_BLUE_LIGHT,
	_FILL_ORANGE_LIGHT,
	_FILL_BLUE_DARK,
	_CENTER,
	_LEFT,
	_RIGHT,
	_BORDER,
	NUM_FMT_RAW,
	NUM_FMT_PCT,
)


def _sheet_budget_vs_actual(wb, heads, fy, month):
	"""Expense Head -> Sub Head -> Item, one row each, Budget/Actual/Util%/
	Variance columns - mirrors the Vue "Budget vs Actual" page's own
	expand-all table exactly (same get_combined_actuals_report tree, same
	Budget=ytd / Actual=total_posted_amt(_ytd) field reads), so the
	exported file always matches what's on screen when "Export Excel" is
	clicked, indentation and grand total included.
	"""
	ws = wb.create_sheet("Budget vs Actual")
	ws.column_dimensions["A"].width = 45
	for col in "BCDE":
		ws.column_dimensions[col].width = 18

	ws["A1"] = f"Budget vs Actual — YTD {month} {fy}"
	ws["A1"].font = Font(bold=True, name="Calibri", size=13)
	ws.merge_cells("A1:E1")

	header_row = 3
	headers = ["Expense Head / Sub Head / Item", "Budget", "Actual", "Util %", "Variance"]
	for col, text in enumerate(headers, start=1):
		c = ws.cell(row=header_row, column=col, value=text)
		c.font, c.fill, c.alignment, c.border = _WHITE_BOLD, _FILL_BLUE_MID, _CENTER, _BORDER

	def item_actual(item):
		return _fv(item.get("total_posted_amt", item.get("total_posted_amt_ytd", 0)))

	def util(budget, actual):
		return round(actual / budget, 4) if budget else 0

	row = header_row + 1
	grand_budget = grand_actual = 0.0

	for head in heads or []:
		head_budget = _fv(head.get("ytd", 0))
		head_actual = _fv(head.get("total_posted_amt_ytd", 0))
		grand_budget += head_budget
		grand_actual += head_actual

		c = ws.cell(row=row, column=1, value=head.get("name", ""))
		c.font, c.fill, c.border, c.alignment = _BOLD, _FILL_BLUE_LIGHT, _BORDER, _LEFT
		for col, val in ((2, head_budget), (3, head_actual)):
			cell = ws.cell(row=row, column=col, value=val)
			cell.number_format, cell.font, cell.fill, cell.border, cell.alignment = NUM_FMT_RAW, _BOLD, _FILL_BLUE_LIGHT, _BORDER, _RIGHT
		u = ws.cell(row=row, column=4, value=util(head_budget, head_actual))
		u.number_format, u.font, u.fill, u.border, u.alignment = NUM_FMT_PCT, _BOLD, _FILL_BLUE_LIGHT, _BORDER, _RIGHT
		v = ws.cell(row=row, column=5, value=head_budget - head_actual)
		v.number_format, v.font, v.fill, v.border, v.alignment = NUM_FMT_RAW, _BOLD, _FILL_BLUE_LIGHT, _BORDER, _RIGHT
		row += 1

		for sub in head.get("sub_heads", []) or []:
			sub_budget = _fv(sub.get("ytd", 0))
			sub_actual = _fv(sub.get("total_posted_amt_ytd", 0))
			c = ws.cell(row=row, column=1, value="  " + sub.get("name", ""))
			c.fill, c.border, c.alignment = _FILL_ORANGE_LIGHT, _BORDER, _LEFT
			for col, val in ((2, sub_budget), (3, sub_actual)):
				cell = ws.cell(row=row, column=col, value=val)
				cell.number_format, cell.fill, cell.border, cell.alignment = NUM_FMT_RAW, _FILL_ORANGE_LIGHT, _BORDER, _RIGHT
			u = ws.cell(row=row, column=4, value=util(sub_budget, sub_actual))
			u.number_format, u.fill, u.border, u.alignment = NUM_FMT_PCT, _FILL_ORANGE_LIGHT, _BORDER, _RIGHT
			v = ws.cell(row=row, column=5, value=sub_budget - sub_actual)
			v.number_format, v.fill, v.border, v.alignment = NUM_FMT_RAW, _FILL_ORANGE_LIGHT, _BORDER, _RIGHT
			row += 1

			for item in sub.get("items", []) or []:
				item_budget = _fv(item.get("ytd", 0))
				item_act = item_actual(item)
				c = ws.cell(row=row, column=1, value="    " + item.get("name", ""))
				c.border, c.alignment = _BORDER, _LEFT
				for col, val in ((2, item_budget), (3, item_act)):
					cell = ws.cell(row=row, column=col, value=val)
					cell.number_format, cell.border, cell.alignment = NUM_FMT_RAW, _BORDER, _RIGHT
				u = ws.cell(row=row, column=4, value=util(item_budget, item_act))
				u.number_format, u.border, u.alignment = NUM_FMT_PCT, _BORDER, _RIGHT
				v = ws.cell(row=row, column=5, value=item_budget - item_act)
				v.number_format, v.border, v.alignment = NUM_FMT_RAW, _BORDER, _RIGHT
				row += 1

		# A head with no sub_heads carries its own items directly (matches
		# the Vue page's tree-walk, which only reads head.items when
		# head.sub_heads is empty).
		if not head.get("sub_heads") and head.get("items"):
			for item in head.get("items", []) or []:
				item_budget = _fv(item.get("ytd", 0))
				item_act = item_actual(item)
				c = ws.cell(row=row, column=1, value="  " + item.get("name", ""))
				c.border, c.alignment = _BORDER, _LEFT
				for col, val in ((2, item_budget), (3, item_act)):
					cell = ws.cell(row=row, column=col, value=val)
					cell.number_format, cell.border, cell.alignment = NUM_FMT_RAW, _BORDER, _RIGHT
				u = ws.cell(row=row, column=4, value=util(item_budget, item_act))
				u.number_format, u.border, u.alignment = NUM_FMT_PCT, _BORDER, _RIGHT
				v = ws.cell(row=row, column=5, value=item_budget - item_act)
				v.number_format, v.border, v.alignment = NUM_FMT_RAW, _BORDER, _RIGHT
				row += 1

	c = ws.cell(row=row, column=1, value="GRAND TOTAL")
	c.font, c.fill, c.border, c.alignment = _WHITE_BOLD, _FILL_BLUE_DARK, _BORDER, _LEFT
	for col, val in ((2, grand_budget), (3, grand_actual)):
		cell = ws.cell(row=row, column=col, value=val)
		cell.number_format, cell.font, cell.fill, cell.border, cell.alignment = NUM_FMT_RAW, _WHITE_BOLD, _FILL_BLUE_DARK, _BORDER, _RIGHT
	u = ws.cell(row=row, column=4, value=util(grand_budget, grand_actual))
	u.number_format, u.font, u.fill, u.border, u.alignment = NUM_FMT_PCT, _WHITE_BOLD, _FILL_BLUE_DARK, _BORDER, _RIGHT
	v = ws.cell(row=row, column=5, value=grand_budget - grand_actual)
	v.number_format, v.font, v.fill, v.border, v.alignment = NUM_FMT_RAW, _WHITE_BOLD, _FILL_BLUE_DARK, _BORDER, _RIGHT


@guest_api
def export_budget_vs_actual(financial_year, month, heads_data):
	fy = financial_year or "2025-26"
	wb = Workbook()
	wb.remove(wb.active)
	_sheet_budget_vs_actual(wb, json.loads(heads_data), fy, month or "March")
	return {"filename": f"Budget_vs_Actual_{fy}_{month or 'March'}.xlsx", "data": _wb_to_b64(wb)}
