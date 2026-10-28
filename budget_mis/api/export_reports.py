import frappe,io
from frappe import _
import frappe
import tempfile
import xlsxwriter
from frappe.utils import nowdate

# * ==============================================================  Import Template Export  =====================================================================================
# @frappe.whitelist()
# def download_finance_budget_import_template(user):

#     if not user:
#         frappe.throw(_("User is required"))

#     # Get Finance User Access
#     doc_name = frappe.db.get_value(
#         "Finance user access",
#         {"user": user},
#         "name"
#     )

#     if not doc_name:
#         frappe.throw(_("No Finance User Access found"))

#     access_doc = frappe.get_doc("Finance user access", doc_name)

#     # Get Import Template ID from Finance User Access
#     if not access_doc.import_template_id:
#         frappe.throw(_("Import Template not linked in Finance User Access"))

#     import_template = frappe.get_doc(
#         "Import Templates",
#         access_doc.import_template_id
#     )

#     template_items = import_template.import_template_item_list

#     if not template_items:
#         frappe.throw(_("No Import Template Items found"))

#     financial_year = frappe.db.get_single_value(
#         "Master Settings",
#         "current_financial_year"
#     )

#     headers = [
#         "Entity / Unit",
#         "Entity / Unit Description",
#         "Cost Center",
#         "Cost Center(Original)",
#         "Cost Center Description",
#         "Location code",
#         "Location code(Original)",
#         "Function / Sub Unit / Division",
#         "State",
#         "Financial year",
#         "Uploaded By",
#         "Type of expense ID (Budget Amounts)",
#         "Head of expense (Budget Amounts)",
#         "Sub head of expense (Budget Amounts)",
#         "Type of expense (Budget Amounts)",
#         "April (Budget Amounts)",
#         "May (Budget Amounts)",
#         "June (Budget Amounts)",
#         "July (Budget Amounts)",
#         "August (Budget Amounts)",
#         "September (Budget Amounts)",
#         "October (Budget Amounts)",
#         "November (Budget Amounts)",
#         "December (Budget Amounts)",
#         "January (Budget Amounts)",
#         "February (Budget Amounts)",
#         "March (Budget Amounts)",
#         "Quarter 1 Total Amount (Budget Amounts)",
#         "Quarter 2 Total Amount (Budget Amounts)",
#         "Quarter 3 Total Amount (Budget Amounts)",
#         "Quarter 4 Total Amount (Budget Amounts)",
#         "Year Total Amount (Budget Amounts)"
#     ]

#     from openpyxl import Workbook
#     from openpyxl.styles import Font, Protection
#     from io import BytesIO
#     import datetime

#     wb = Workbook()
#     ws = wb.active
#     ws.title = "Finance Budget Import"

#     ws.append(headers)

#     for cell in ws[1]:
#         cell.font = Font(bold=True)

#     ws.freeze_panes = "A2"

#     row_index = 2

#     for mapping in access_doc.mapping:

#         first_row = True

#         for item in template_items:

#             if first_row:
#                 parent_values = [
#                     mapping.unit,
#                     mapping.unit_description,
#                     mapping.cost_center,
#                     mapping.cost_center_erp,
#                     mapping.cost_center_description,
#                     mapping.location_code,
#                     mapping.location_code_erp,
#                     mapping.location_description,
#                     mapping.state,
#                     financial_year,
#                     user
#                 ]
#                 first_row = False
#             else:
#                 parent_values = [""] * 11

#             ws.append(parent_values + [
#                 item.type_of_expense_id,
#                 item.head_of_expense,
#                 item.sub_head_of_expense,
#                 item.type_of_expense,
#                 0.00, 0.00, 0.00,
#                 0.00, 0.00, 0.00,
#                 0.00, 0.00, 0.00,
#                 0.00, 0.00, 0.00,
#                 0.00, 0.00, 0.00, 0.00, 0.00
#             ])

#             # Quarter formulas
#             ws[f"AB{row_index}"] = f"=SUM(P{row_index}:R{row_index})"
#             ws[f"AC{row_index}"] = f"=SUM(S{row_index}:U{row_index})"
#             ws[f"AD{row_index}"] = f"=SUM(V{row_index}:X{row_index})"
#             ws[f"AE{row_index}"] = f"=SUM(Y{row_index}:AA{row_index})"

#             # Year total = Sum of quarters
#             ws[f"AF{row_index}"] = f"=SUM(AB{row_index}:AE{row_index})"

#             # Format numbers
#             for col in range(16, 33):
#                 ws.cell(row=row_index, column=col).number_format = '0.00'

#             row_index += 1

#     # Lock everything
#     for row in ws.iter_rows(min_row=1, max_row=row_index - 1):
#         for cell in row:
#             cell.protection = Protection(locked=True)

#     # Unlock only monthly columns (April–March)
#     for row in ws.iter_rows(min_row=2, max_row=row_index - 1, min_col=16, max_col=27):
#         for cell in row:
#             cell.protection = Protection(locked=False)

#     ws.protection.sheet = True
#     ws.protection.password = "[removed]"

#     # Auto column width
#     for column in ws.columns:
#         max_length = 0
#         column_letter = column[0].column_letter
#         for cell in column:
#             if cell.value:
#                 max_length = max(max_length, len(str(cell.value)))
#         ws.column_dimensions[column_letter].width = max_length + 2

#     output = BytesIO()
#     wb.save(output)
#     output.seek(0)

#     current_date = datetime.datetime.now().strftime("%Y-%m-%d")
#     frappe.response["filename"] = f"Budget_mis_Import_{current_date}.xlsx"
#     frappe.response["filecontent"] = output.getvalue()
#     frappe.response["type"] = "download"

# @frappe.whitelist()
# def download_finance_budget_import_template(user):

#     import frappe
#     from frappe import _
#     import datetime
#     from io import BytesIO
#     import xlsxwriter

#     if not user:
#         frappe.throw(_("User is required"))

#     # Get Finance User Access
#     doc_name = frappe.db.get_value(
#         "Finance user access",
#         {"user": user},
#         "name"
#     )

#     if not doc_name:
#         frappe.throw(_("No Finance User Access found"))

#     access_doc = frappe.get_doc("Finance user access", doc_name)

#     if not access_doc.import_template_id:
#         frappe.throw(_("Import Template not linked in Finance User Access"))

#     import_template = frappe.get_doc(
#         "Import Templates",
#         access_doc.import_template_id
#     )

#     template_items = import_template.import_template_item_list

#     if not template_items:
#         frappe.throw(_("No Import Template Items found"))

#     financial_year = frappe.db.get_single_value(
#         "Master Settings",
#         "current_financial_year"
#     )

#     headers = [
#         "Entity / Unit",
#         "Entity / Unit Description",
#         "Cost Center",
#         "Cost Center(Original)",
#         "Cost Center Description",
#         "Location code",
#         "Location code(Original)",
#         "Function / Sub Unit / Division",
#         "State",
#         "Financial year",
#         "Uploaded By",
#         "Type of expense ID (Budget Amounts)",
#         "Head of expense (Budget Amounts)",
#         "Sub head of expense (Budget Amounts)",
#         "Type of expense (Budget Amounts)",
#         "April (Budget Amounts)",
#         "May (Budget Amounts)",
#         "June (Budget Amounts)",
#         "July (Budget Amounts)",
#         "August (Budget Amounts)",
#         "September (Budget Amounts)",
#         "October (Budget Amounts)",
#         "November (Budget Amounts)",
#         "December (Budget Amounts)",
#         "January (Budget Amounts)",
#         "February (Budget Amounts)",
#         "March (Budget Amounts)",
#         "Quarter 1 Total Amount (Budget Amounts)",
#         "Quarter 2 Total Amount (Budget Amounts)",
#         "Quarter 3 Total Amount (Budget Amounts)",
#         "Quarter 4 Total Amount (Budget Amounts)",
#         "Year Total Amount (Budget Amounts)"
#     ]

#     output = BytesIO()

#     workbook = xlsxwriter.Workbook(output)
#     worksheet = workbook.add_worksheet("Finance Budget Import")

#     # Formats
#     header_format = workbook.add_format({
#         "bold": True,
#         "locked": True
#     })

#     locked_text = workbook.add_format({
#         "locked": True
#     })

#     locked_number = workbook.add_format({
#         "locked": True,
#         "num_format": "0.00"
#     })

#     unlocked_number = workbook.add_format({
#         "locked": False,
#         "num_format": "0.00"
#     })

#     # Write headers
#     for col, header in enumerate(headers):
#         worksheet.write(0, col, header, header_format)
#         worksheet.set_column(col, col, 22)

#     worksheet.freeze_panes(1, 0)

#     row_index = 1

#     for mapping in access_doc.mapping:

#         first_row = True

#         for item in template_items:

#             if first_row:
#                 parent_values = [
#                     mapping.unit,
#                     mapping.unit_description,
#                     mapping.cost_center,
#                     mapping.cost_center_erp,
#                     mapping.cost_center_description,
#                     mapping.location_code,
#                     mapping.location_code_erp,
#                     mapping.location_description,
#                     mapping.state,
#                     financial_year,
#                     user
#                 ]
#                 first_row = False
#             else:
#                 parent_values = [""] * 11

#             row_data = parent_values + [
#                 item.type_of_expense_id,
#                 item.head_of_expense,
#                 item.sub_head_of_expense,
#                 item.type_of_expense
#             ]

#             # Write locked columns
#             for col, value in enumerate(row_data):
#                 worksheet.write(row_index, col, value, locked_text)

#             # Write editable month columns (April–March)
#             for col in range(15, 27):
#                 worksheet.write(row_index, col, 0.00, unlocked_number)

#             # Quarter formulas
#             r = row_index + 1

#             worksheet.write_formula(row_index, 27, f"=SUM(P{r}:R{r})", locked_number)
#             worksheet.write_formula(row_index, 28, f"=SUM(S{r}:U{r})", locked_number)
#             worksheet.write_formula(row_index, 29, f"=SUM(V{r}:X{r})", locked_number)
#             worksheet.write_formula(row_index, 30, f"=SUM(Y{r}:AA{r})", locked_number)

#             # Year total
#             worksheet.write_formula(row_index, 31, f"=SUM(AB{r}:AE{r})", locked_number)

#             row_index += 1

#     # Enable sheet protection
#     worksheet.protect("[REDACTED-PASSWORD]")

#     workbook.close()

#     output.seek(0)

#     current_date = datetime.datetime.now().strftime("%Y-%m-%d")

#     frappe.response["filename"] = f"Budget_mis_Import_{current_date}.xlsx"
#     frappe.response["filecontent"] = output.getvalue()
#     frappe.response["type"] = "download"


# @frappe.whitelist()
# def start_budget_template_generation(user):

#     if not user:
#         frappe.throw("User required")

#     frappe.enqueue(
#         "budget_mis.api.export_reports.generate_budget_template",
#         queue="long",
#         timeout=1800,
#         user=user
#     )

#     return {"status": "started"}


# def generate_budget_template(user):

#     doc_name = frappe.db.get_value(
#         "Finance user access",
#         {"user": user},
#         "name"
#     )

#     if not doc_name:
#         return

#     access_doc = frappe.get_doc("Finance user access", doc_name)

#     import_template = frappe.get_doc(
#         "Import Templates",
#         access_doc.import_template_id
#     )

#     template_items = import_template.import_template_item_list

#     financial_year = frappe.db.get_single_value(
#         "Master Settings",
#         "current_financial_year"
#     )

#     tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")

#     workbook = xlsxwriter.Workbook(tmp.name)
#     worksheet = workbook.add_worksheet("Finance Budget Import")

#     header_format = workbook.add_format({"bold": True, "locked": True})
#     locked = workbook.add_format({"locked": True})
#     unlocked = workbook.add_format({"locked": False, "num_format": "0.00"})

#     headers = [
#         "Entity / Unit",
#         "Entity / Unit Description",
#         "Cost Center",
#         "Cost Center(Original)",
#         "Cost Center Description",
#         "Location code",
#         "Location code(Original)",
#         "Function / Sub Unit / Division",
#         "State",
#         "Financial year",
#         "Uploaded By",
#         "Type of expense ID (Budget Amounts)",
#         "Head of expense (Budget Amounts)",
#         "Sub head of expense (Budget Amounts)",
#         "Type of expense (Budget Amounts)",
#         "April (Budget Amounts)",
#         "May (Budget Amounts)",
#         "June (Budget Amounts)",
#         "July (Budget Amounts)",
#         "August (Budget Amounts)",
#         "September (Budget Amounts)",
#         "October (Budget Amounts)",
#         "November (Budget Amounts)",
#         "December (Budget Amounts)",
#         "January (Budget Amounts)",
#         "February (Budget Amounts)",
#         "March (Budget Amounts)",
#         "Quarter 1 Total Amount (Budget Amounts)",
#         "Quarter 2 Total Amount (Budget Amounts)",
#         "Quarter 3 Total Amount (Budget Amounts)",
#         "Quarter 4 Total Amount (Budget Amounts)",
#         "Year Total Amount (Budget Amounts)"
#     ]

#     # Track column width
#     col_widths = [len(h) for h in headers]

#     for col, header in enumerate(headers):
#         worksheet.write(0, col, header, header_format)

#     worksheet.freeze_panes(1, 0)

#     row_index = 1

#     for mapping in access_doc.mapping:

#         first_row = True

#         for item in template_items:

#             if first_row:
#                 parent_values = [
#                     mapping.unit,
#                     mapping.unit_description,
#                     mapping.cost_center,
#                     mapping.cost_center_erp,
#                     mapping.cost_center_description,
#                     mapping.location_code,
#                     mapping.location_code_erp,
#                     mapping.location_description,
#                     mapping.state,
#                     financial_year,
#                     user
#                 ]
#                 first_row = False
#             else:
#                 parent_values = [""] * 11

#             row = parent_values + [
#                 item.type_of_expense_id,
#                 item.head_of_expense,
#                 item.sub_head_of_expense,
#                 item.type_of_expense
#             ]

#             for col, val in enumerate(row):

#                 worksheet.write(row_index, col, val, locked)

#                 if val:
#                     col_widths[col] = max(col_widths[col], len(str(val)))

#             # Editable months
#             for col in range(15, 27):
#                 worksheet.write(row_index, col, 0, unlocked)

#             r = row_index + 1

#             worksheet.write_formula(row_index, 27, f"=SUM(P{r}:R{r})")
#             worksheet.write_formula(row_index, 28, f"=SUM(S{r}:U{r})")
#             worksheet.write_formula(row_index, 29, f"=SUM(V{r}:X{r})")
#             worksheet.write_formula(row_index, 30, f"=SUM(Y{r}:AA{r})")
#             worksheet.write_formula(row_index, 31, f"=SUM(AB{r}:AE{r})")

#             row_index += 1

#     worksheet.protect("[REDACTED-PASSWORD]")

#     # Apply calculated widths
#     for i, width in enumerate(col_widths):
#         worksheet.set_column(i, i, width + 3)

#     workbook.close()

#     frappe.cache().set_value(
#         f"budget_template_{user}",
#         tmp.name,
#         expires_in_sec=3600
#     )


# @frappe.whitelist()
# def download_generated_template(user):

#     path = frappe.cache().get_value(f"budget_template_{user}")

#     if not path:
#         return {"status": "processing"}

#     with open(path, "rb") as f:

#         frappe.response["filename"] = f"Budget_mis_Import_{nowdate()}.xlsx"
#         frappe.response["filecontent"] = f.read()
#         frappe.response["type"] = "download"



    # import frappe
    # import tempfile
    # import xlsxwriter
    # from frappe.utils import nowdate


    # @frappe.whitelist()
    # def start_budget_template_generation(user, entity_data=None):

    #     if not user:
    #         frappe.throw("User required")

    #     frappe.enqueue(
    #         "budget_mis.api.export_reports.generate_budget_template",
    #         queue="long",
    #         timeout=1800,
    #         user=user,
    #         entity_data=entity_data
    #     )

    #     return {"status": "started"}


    # def generate_budget_template(user, entity_data=None):

    #     # ── Parse entity_data ─────────────────────────────────
    #     # Can be:
    #     #   None          → use all mappings from access_doc
    #     #   JSON string   → parse into list or single dict
    #     #   list          → multiple selected rows
    #     #   dict          → single selected row (legacy)

    #     if entity_data and isinstance(entity_data, str):
    #         entity_data = frappe.parse_json(entity_data)

    #     # Normalise to a list (or None)
    #     if isinstance(entity_data, dict):
    #         entity_data = [entity_data]   # wrap single row in a list

    #     # entity_data is now either a list of dicts or None

    #     # ── Fetch Finance User Access ──────────────────────────
    #     doc_name = frappe.db.get_value(
    #         "Finance user access",
    #         {"user": user},
    #         "name"
    #     )

    #     access_doc = frappe.get_doc("Finance user access", doc_name) if doc_name else None

    #     if not access_doc:
    #         return

    #     # ── Import Template ────────────────────────────────────
    #     import_template = frappe.get_doc(
    #         "Import Templates",
    #         access_doc.import_template_id
    #     )

    #     template_items = import_template.import_template_item_list

    #     financial_year = frappe.db.get_single_value(
    #         "Master Settings",
    #         "current_financial_year"
    #     )

    #     # ── Decide mappings source ─────────────────────────────
    #     # If entity_data list provided → use those rows
    #     # Otherwise → use all rows from access_doc.mapping
    #     if entity_data:
    #         mappings       = entity_data          # list of plain dicts from frontend
    #         use_dict       = True                 # access via .get()
    #     else:
    #         mappings       = access_doc.mapping   # list of child doc objects
    #         use_dict       = False                # access via getattr()

    #     # ── Helper: read a field from either a dict or a child doc ──
    #     def get_field(field, mapping):
    #         if use_dict:
    #             return mapping.get(field) or ""
    #         return getattr(mapping, field, "") or ""

    #     # ── Create Excel ───────────────────────────────────────
    #     tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")

    #     workbook  = xlsxwriter.Workbook(tmp.name)
    #     worksheet = workbook.add_worksheet("Finance Budget Import")

    #     header_format = workbook.add_format({"bold": True, "locked": True})
    #     locked        = workbook.add_format({"locked": True})
    #     unlocked      = workbook.add_format({"locked": False, "num_format": "0.00"})

    #     headers = [
    #         "Entity / Unit",
    #         "Entity / Unit Description",
    #         "Cost Center",
    #         "Cost Center(Original)",
    #         "Cost Center Description",
    #         "Location code",
    #         "Location code(Original)",
    #         "Function / Sub Unit / Division",
    #         "State",
    #         "Financial year",
    #         "Uploaded By",
    #         "Type of expense ID (Budget Amounts)",
    #         "Head of expense (Budget Amounts)",
    #         "Sub head of expense (Budget Amounts)",
    #         "Type of expense (Budget Amounts)",
    #         "April (Budget Amounts)",
    #         "May (Budget Amounts)",
    #         "June (Budget Amounts)",
    #         "July (Budget Amounts)",
    #         "August (Budget Amounts)",
    #         "September (Budget Amounts)",
    #         "October (Budget Amounts)",
    #         "November (Budget Amounts)",
    #         "December (Budget Amounts)",
    #         "January (Budget Amounts)",
    #         "February (Budget Amounts)",
    #         "March (Budget Amounts)",
    #         "Quarter 1 Total Amount (Budget Amounts)",
    #         "Quarter 2 Total Amount (Budget Amounts)",
    #         "Quarter 3 Total Amount (Budget Amounts)",
    #         "Quarter 4 Total Amount (Budget Amounts)",
    #         "Year Total Amount (Budget Amounts)"
    #     ]

    #     col_widths = [len(h) for h in headers]

    #     for col, header in enumerate(headers):
    #         worksheet.write(0, col, header, header_format)

    #     worksheet.freeze_panes(1, 0)

    #     row_index = 1

    #     # ── Write one block of rows per mapping ───────────────
    #     for mapping in mappings:

    #         first_row = True

    #         for item in template_items:

    #             if first_row:
    #                 parent_values = [
    #                     get_field("unit",                   mapping),
    #                     get_field("unit_description",       mapping),
    #                     get_field("cost_center",            mapping),
    #                     get_field("cost_center_erp",        mapping),
    #                     get_field("cost_center_description",mapping),
    #                     get_field("location_code",          mapping),
    #                     get_field("location_code_erp",      mapping),
    #                     get_field("location_description",   mapping),
    #                     get_field("state",                  mapping),
    #                     financial_year,
    #                     user
    #                 ]
    #                 first_row = False
    #             else:
    #                 parent_values = [""] * 11

    #             row = parent_values + [
    #                 item.type_of_expense_id,
    #                 item.head_of_expense,
    #                 item.sub_head_of_expense,
    #                 item.type_of_expense
    #             ]

    #             for col, val in enumerate(row):
    #                 worksheet.write(row_index, col, val, locked)
    #                 if val:
    #                     col_widths[col] = max(col_widths[col], len(str(val)))

    #             for col in range(15, 27):
    #                 worksheet.write(row_index, col, 0, unlocked)

    #             r = row_index + 1
    #             worksheet.write_formula(row_index, 27, f"=SUM(P{r}:R{r})")
    #             worksheet.write_formula(row_index, 28, f"=SUM(S{r}:U{r})")
    #             worksheet.write_formula(row_index, 29, f"=SUM(V{r}:X{r})")
    #             worksheet.write_formula(row_index, 30, f"=SUM(Y{r}:AA{r})")
    #             worksheet.write_formula(row_index, 31, f"=SUM(AB{r}:AE{r})")

    #             row_index += 1

    #     worksheet.protect("[REDACTED-PASSWORD]")

    #     for i, width in enumerate(col_widths):
    #         worksheet.set_column(i, i, width + 3)

    #     workbook.close()

    #     frappe.cache().set_value(
    #         f"budget_template_{user}",
    #         tmp.name,
    #         expires_in_sec=3600
    #     )


    # @frappe.whitelist()
    # def download_generated_template(user):

    #     path = frappe.cache().get_value(f"budget_template_{user}")

    #     if not path:
    #         return {"status": "processing"}

    #     with open(path, "rb") as f:
    #         frappe.response["filename"]    = f"Budget_mis_Import_{nowdate()}.xlsx"
    #         frappe.response["filecontent"] = f.read()
    #         frappe.response["type"]        = "download"




import frappe
import tempfile
import xlsxwriter
from frappe.utils import nowdate


@frappe.whitelist()
def start_budget_template_generation(entity_data=None):

    user = frappe.session.user

    frappe.enqueue(
        "budget_mis.api.export_reports.generate_budget_template",
        queue="long",
        timeout=1800,
        user=user,
        entity_data=entity_data
    )

    return {"status": "started"}


def generate_budget_template(user=None, entity_data=None):

    user = user or frappe.session.user

    # ── Parse entity_data ─────────────────────────────────
    if entity_data and isinstance(entity_data, str):
        entity_data = frappe.parse_json(entity_data)

    if isinstance(entity_data, dict):
        entity_data = [entity_data]

    # ── Fetch Finance User Access ─────────────────────────
    doc_name = frappe.db.get_value(
        "Finance user access",
        {"user": user},
        "name"
    )

    access_doc = frappe.get_doc("Finance user access", doc_name) if doc_name else None

    if not access_doc:
        return

    allow_edit = access_doc.allow_edit_template or 0

    # ── Import Template ───────────────────────────────────
    import_template = frappe.get_doc(
        "Import Templates",
        access_doc.import_template_id
    )

    template_items = import_template.import_template_item_list

    financial_year = frappe.db.get_single_value(
        "Master Settings",
        "current_financial_year"
    )

    # ── Decide mappings source ─────────────────────────────
    if entity_data:
        mappings = entity_data
        use_dict = True
    else:
        mappings = access_doc.mapping
        use_dict = False

    # ── Helper ────────────────────────────────────────────
    def get_field(frontend_key, child_key, mapping):
        if use_dict:
            return mapping.get(frontend_key) or ""
        return getattr(mapping, child_key, "") or ""

    # ── Create Excel ──────────────────────────────────────
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")

    workbook  = xlsxwriter.Workbook(tmp.name)
    worksheet = workbook.add_worksheet("Finance Budget Import")

    header_format = workbook.add_format({"bold": True, "locked": True})

    if allow_edit:
        locked = workbook.add_format({"locked": False})
        unlocked = workbook.add_format({"locked": False, "num_format": "0.00"})
    else:
        locked = workbook.add_format({"locked": True})
        unlocked = workbook.add_format({"locked": False, "num_format": "0.00"})

    headers = [
        "Entity / Unit",
        "Entity / Unit Description",
        "Cost Center",
        "Cost Center(Original)",
        "Cost Center Description",
        "Location code",
        "Location code(Original)",
        "Location Description",
        "State",
        "Financial year",
        "Uploaded By",
        "Type of expense ID (Budget Amounts)",
        "Head of expense (Budget Amounts)",
        "Sub head of expense (Budget Amounts)",
        "Type of expense (Budget Amounts)",
        "April (Budget Amounts)",
        "May (Budget Amounts)",
        "June (Budget Amounts)",
        "July (Budget Amounts)",
        "August (Budget Amounts)",
        "September (Budget Amounts)",
        "October (Budget Amounts)",
        "November (Budget Amounts)",
        "December (Budget Amounts)",
        "January (Budget Amounts)",
        "February (Budget Amounts)",
        "March (Budget Amounts)",
        "Quarter 1 Total Amount (Budget Amounts)",
        "Quarter 2 Total Amount (Budget Amounts)",
        "Quarter 3 Total Amount (Budget Amounts)",
        "Quarter 4 Total Amount (Budget Amounts)",
        "Year Total Amount (Budget Amounts)"
    ]

    col_widths = [len(h) for h in headers]

    for col, header in enumerate(headers):
        worksheet.write(0, col, header, header_format)

    worksheet.freeze_panes(1, 0)

    row_index = 1

    for mapping in mappings:

        first_row = True

        for item in template_items:

            if first_row:
                parent_values = [
                    get_field("unit", "unit", mapping),
                    get_field("unit_description", "unit_description", mapping),

                    # ✅ FIXED COST CENTER
                    get_field("cost_center_id", "cost_center", mapping),
                    get_field("cost_center", "cost_center_erp", mapping),

                    get_field("cost_center_description", "cost_center_description", mapping),

                    # ✅ FIXED LOCATION
                    get_field("location_code_id", "location_code", mapping),
                    get_field("location_code", "location_code_erp", mapping),

                    get_field("location_description", "location_description", mapping),
                    get_field("state", "state", mapping),
                    financial_year,
                    user
                ]
                first_row = False
            else:
                parent_values = [""] * 11

            row = parent_values + [
                item.type_of_expense_id,
                item.head_of_expense,
                item.sub_head_of_expense,
                item.type_of_expense
            ]

            for col, val in enumerate(row):
                worksheet.write(row_index, col, val, locked)
                if val:
                    col_widths[col] = max(col_widths[col], len(str(val)))

            # Editable monthly fields
            for col in range(15, 27):
                worksheet.write(row_index, col, 0, unlocked)

            r = row_index + 1
            worksheet.write_formula(row_index, 27, f"=SUM(P{r}:R{r})")
            worksheet.write_formula(row_index, 28, f"=SUM(S{r}:U{r})")
            worksheet.write_formula(row_index, 29, f"=SUM(V{r}:X{r})")
            worksheet.write_formula(row_index, 30, f"=SUM(Y{r}:AA{r})")
            worksheet.write_formula(row_index, 31, f"=SUM(AB{r}:AE{r})")

            row_index += 1

    if not allow_edit:
        worksheet.protect("[REDACTED-PASSWORD]")

    for i, width in enumerate(col_widths):
        worksheet.set_column(i, i, width + 3)

    workbook.close()

    frappe.cache().set_value(
        f"budget_template_{user}",
        tmp.name,
        expires_in_sec=3600
    )


@frappe.whitelist()
def download_generated_template():

    user = frappe.session.user

    path = frappe.cache().get_value(f"budget_template_{user}")

    if not path:
        return {"status": "processing"}

    with open(path, "rb") as f:
        frappe.response["filename"]    = f"Budget_mis_Import_{nowdate()}.xlsx"
        frappe.response["filecontent"] = f.read()
        frappe.response["type"]        = "download"
        



# import tempfile
# import xlsxwriter
# from frappe.utils import nowdate
# import frappe


# @frappe.whitelist()
# def download_finance_budget_import_template(user):

#     if not user:
#         frappe.throw("User required")

#     doc_name = frappe.db.get_value(
#         "Finance user access",
#         {"user": user},
#         "name"
#     )

#     if not doc_name:
#         frappe.throw("No Finance User Access record found for this user.")

#     access_doc = frappe.get_doc("Finance user access", doc_name)

#     # ── allow_edit_template checkbox ──────────────────────
#     # Checked (1)   → skip sheet protection → entire sheet is editable
#     # Unchecked (0) → protect sheet         → only month columns are editable
#     is_editable = access_doc.allow_edit_template == 1

#     import_template = frappe.get_doc(
#         "Import Templates",
#         access_doc.import_template_id
#     )

#     template_items = import_template.import_template_item_list

#     financial_year = frappe.db.get_single_value(
#         "Master Settings",
#         "current_financial_year"
#     )

#     tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")

#     workbook  = xlsxwriter.Workbook(tmp.name)
#     worksheet = workbook.add_worksheet("Finance Budget Import")

#     header_format = workbook.add_format({"bold": True, "locked": True})
#     locked        = workbook.add_format({"locked": True})

#     # Month columns are always written with locked=False.
#     # When is_editable=False → sheet is protected → only these cells are editable.
#     # When is_editable=True  → sheet is not protected → all cells are editable anyway.
#     unlocked = workbook.add_format({"locked": False, "num_format": "0.00"})

#     headers = [
#         "Entity / Unit",
#         "Entity / Unit Description",
#         "Cost Center",
#         "Cost Center(Original)",
#         "Cost Center Description",
#         "Location code",
#         "Location code(Original)",
#         "Function / Sub Unit / Division",
#         "State",
#         "Financial year",
#         "Uploaded By",
#         "Type of expense ID (Budget Amounts)",
#         "Head of expense (Budget Amounts)",
#         "Sub head of expense (Budget Amounts)",
#         "Type of expense (Budget Amounts)",
#         "April (Budget Amounts)",
#         "May (Budget Amounts)",
#         "June (Budget Amounts)",
#         "July (Budget Amounts)",
#         "August (Budget Amounts)",
#         "September (Budget Amounts)",
#         "October (Budget Amounts)",
#         "November (Budget Amounts)",
#         "December (Budget Amounts)",
#         "January (Budget Amounts)",
#         "February (Budget Amounts)",
#         "March (Budget Amounts)",
#         "Quarter 1 Total Amount (Budget Amounts)",
#         "Quarter 2 Total Amount (Budget Amounts)",
#         "Quarter 3 Total Amount (Budget Amounts)",
#         "Quarter 4 Total Amount (Budget Amounts)",
#         "Year Total Amount (Budget Amounts)"
#     ]

#     col_widths = [len(h) for h in headers]

#     for col, header in enumerate(headers):
#         worksheet.write(0, col, header, header_format)

#     worksheet.freeze_panes(1, 0)

#     row_index = 1

#     for mapping in access_doc.mapping:

#         first_row = True

#         for item in template_items:

#             if first_row:
#                 parent_values = [
#                     mapping.unit,
#                     mapping.unit_description,
#                     mapping.cost_center,
#                     mapping.cost_center_erp,
#                     mapping.cost_center_description,
#                     mapping.location_code,
#                     mapping.location_code_erp,
#                     mapping.location_description,
#                     mapping.state,
#                     financial_year,
#                     user
#                 ]
#                 first_row = False
#             else:
#                 parent_values = [""] * 11

#             row = parent_values + [
#                 item.type_of_expense_id,
#                 item.head_of_expense,
#                 item.sub_head_of_expense,
#                 item.type_of_expense
#             ]

#             for col, val in enumerate(row):
#                 worksheet.write(row_index, col, val, locked)
#                 if val:
#                     col_widths[col] = max(col_widths[col], len(str(val)))

#             # Month columns (April–March)
#             for col in range(15, 27):
#                 worksheet.write(row_index, col, 0, unlocked)

#             r = row_index + 1

#             worksheet.write_formula(row_index, 27, f"=SUM(P{r}:R{r})")
#             worksheet.write_formula(row_index, 28, f"=SUM(S{r}:U{r})")
#             worksheet.write_formula(row_index, 29, f"=SUM(V{r}:X{r})")
#             worksheet.write_formula(row_index, 30, f"=SUM(Y{r}:AA{r})")
#             worksheet.write_formula(row_index, 31, f"=SUM(AB{r}:AE{r})")

#             row_index += 1

#     # Protect only when allow_edit_template is unchecked
#     if not is_editable:
#         worksheet.protect("[REDACTED-PASSWORD]")

#     for i, width in enumerate(col_widths):
#         worksheet.set_column(i, i, width + 3)

#     workbook.close()

#     with open(tmp.name, "rb") as f:
#         frappe.response["filename"]    = f"Budget_Import_{nowdate()}.xlsx"
#         frappe.response["filecontent"] = f.read()
#         frappe.response["type"]        = "download"






# import frappe
# import tempfile
# import xlsxwriter
# from frappe.utils import nowdate, cint


# # ✅ START JOB (with lock)
# @frappe.whitelist()
# def start_budget_template_generation(user):

#     if not user:
#         frappe.throw("User required")

#     job_key = f"budget_template_job_{user}"

#     # Prevent duplicate jobs
#     if frappe.cache().get_value(job_key):
#         return {"status": "already_running"}

#     frappe.cache().set_value(job_key, True, expires_in_sec=600)

#     frappe.enqueue(
#         "budget_mis.api.export_reports.generate_budget_template",
#         queue="long",
#         timeout=1800,
#         user=user
#     )

#     return {"status": "started"}


# # ✅ BACKGROUND JOB
# def generate_budget_template(user):

#     try:
#         frappe.cache().delete_value(f"budget_template_{user}")

#         doc_name = frappe.db.get_value(
#             "Finance user access",
#             {"user": user},
#             "name"
#         )

#         if not doc_name:
#             return

#         access_doc = frappe.get_doc("Finance user access", doc_name)

#         # ✅ FIX checkbox inconsistency
#         is_editable = cint(access_doc.allow_edit_template) == 1

#         import_template = frappe.get_doc(
#             "Import Templates",
#             access_doc.import_template_id
#         )

#         template_items = import_template.import_template_item_list

#         financial_year = frappe.db.get_single_value(
#             "Master Settings",
#             "current_financial_year"
#         )

#         tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")

#         # ✅ Faster workbook
#         workbook = xlsxwriter.Workbook(tmp.name, {'constant_memory': True})
#         worksheet = workbook.add_worksheet("Finance Budget Import")

#         # Formats
#         unlocked_format = workbook.add_format({
#             "locked": False,
#             "num_format": "0.00"
#         })

#         locked_format = workbook.add_format({"locked": True})

#         header_format = workbook.add_format({
#             "bold": True,
#             "locked": not is_editable
#         })

#         headers = [
#             "Entity / Unit", "Entity / Unit Description", "Cost Center",
#             "Cost Center(Original)", "Cost Center Description",
#             "Location code", "Location code(Original)",
#             "Function / Sub Unit / Division", "State",
#             "Financial year", "Uploaded By",
#             "Type of expense ID (Budget Amounts)",
#             "Head of expense (Budget Amounts)",
#             "Sub head of expense (Budget Amounts)",
#             "Type of expense (Budget Amounts)",
#             "April (Budget Amounts)", "May (Budget Amounts)",
#             "June (Budget Amounts)", "July (Budget Amounts)",
#             "August (Budget Amounts)", "September (Budget Amounts)",
#             "October (Budget Amounts)", "November (Budget Amounts)",
#             "December (Budget Amounts)", "January (Budget Amounts)",
#             "February (Budget Amounts)", "March (Budget Amounts)",
#             "Quarter 1 Total Amount (Budget Amounts)",
#             "Quarter 2 Total Amount (Budget Amounts)",
#             "Quarter 3 Total Amount (Budget Amounts)",
#             "Quarter 4 Total Amount (Budget Amounts)",
#             "Year Total Amount (Budget Amounts)"
#         ]

#         # Write headers
#         for col, header in enumerate(headers):
#             worksheet.write(0, col, header, header_format)

#         worksheet.freeze_panes(1, 0)

#         row_index = 1

#         base_format = unlocked_format if is_editable else locked_format

#         for mapping in access_doc.mapping:

#             first_row = True

#             for item in template_items:

#                 if first_row:
#                     parent_values = [
#                         mapping.unit,
#                         mapping.unit_description,
#                         mapping.cost_center,
#                         mapping.cost_center_erp,
#                         mapping.cost_center_description,
#                         mapping.location_code,
#                         mapping.location_code_erp,
#                         mapping.location_description,
#                         mapping.state,
#                         financial_year,
#                         user
#                     ]
#                     first_row = False
#                 else:
#                     parent_values = [""] * 11

#                 row = parent_values + [
#                     item.type_of_expense_id,
#                     item.head_of_expense,
#                     item.sub_head_of_expense,
#                     item.type_of_expense
#                 ]

#                 # 🚀 Fast row write
#                 worksheet.write_row(row_index, 0, row, base_format)

#                 # Month columns always editable
#                 for col in range(15, 27):
#                     worksheet.write(row_index, col, 0, unlocked_format)

#                 r = row_index + 1

#                 worksheet.write_formula(row_index, 27, f"=SUM(P{r}:R{r})")
#                 worksheet.write_formula(row_index, 28, f"=SUM(S{r}:U{r})")
#                 worksheet.write_formula(row_index, 29, f"=SUM(V{r}:X{r})")
#                 worksheet.write_formula(row_index, 30, f"=SUM(Y{r}:AA{r})")
#                 worksheet.write_formula(row_index, 31, f"=SUM(AB{r}:AE{r})")

#                 row_index += 1

#         # ✅ Apply filter ONLY when editable
#         if is_editable and row_index > 1:
#             worksheet.autofilter(0, 0, row_index - 1, len(headers) - 1)

#         # Protect if needed
#         if not is_editable:
#             worksheet.protect("[REDACTED-PASSWORD]")

#         worksheet.set_column(0, len(headers) - 1, 22)

#         workbook.close()

#         frappe.cache().set_value(
#             f"budget_template_{user}",
#             tmp.name,
#             expires_in_sec=3600
#         )

#     finally:
#         # Always release lock
#         frappe.cache().delete_value(f"budget_template_job_{user}")


# # ✅ STATUS API
# @frappe.whitelist()
# def check_template_status(user):

#     path = frappe.cache().get_value(f"budget_template_{user}")

#     if path:
#         return {"status": "ready"}

#     return {"status": "processing"}


# # ✅ DOWNLOAD API
# @frappe.whitelist()
# def download_generated_template(user):

#     path = frappe.cache().get_value(f"budget_template_{user}")

#     if not path:
#         return {"status": "processing"}

#     with open(path, "rb") as f:
#         frappe.response["filename"] = f"Budget_mis_Import_{nowdate()}.xlsx"
#         frappe.response["filecontent"] = f.read()
#         frappe.response["type"] = "download"





















# # ═══════════════════════════════════════════════════════════════════════════════
# # SHEET: Budget & Estimate
# # ═══════════════════════════════════════════════════════════════════════════════

def _sheet_be(wb, sheet_name, be_data, fy, plan_label, est_label,
              org_name="Azim Premji Foundation"):
    ws = wb.create_sheet(title=sheet_name[:31])
    entities = be_data or []
    n        = len(entities)
    thin = Side(style="thin")
    brd  = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Col layout (1-based): A unused, B=HEAD, C=TYPE, then 2*n entity cols, then 2 grand cols
    c_head     = 2
    c_type     = 3
    c_ent_start= 4              # first entity Plan col
    c_grand_p  = c_ent_start + n * 2       # Grand Total Plan
    c_grand_e  = c_grand_p + 1            # Grand Total Est
    last_col   = c_grand_e

    PF   = "ytd"
    EF   = "total_posted_amt_ytd"
    IF_E = "total_posted_amt"

    def _sec_v(entity, sname, f):
        for s in entity.get("actuals", []):
            if _is_grand_total_section(s): continue
            if s.get("name") == sname:
                return float(s.get(f, 0) or 0)
        return 0.0

    def _sub_v(entity, sname, subname, f):
        for s in entity.get("actuals", []):
            if _is_grand_total_section(s): continue
            if s.get("name") == sname:
                for sub in s.get("sub_heads", []):
                    if sub.get("name") == subname:
                        return float(sub.get(f, 0) or 0)
        return 0.0

    def _item_v(entity, iname, use_est=False):
        f = IF_E if use_est else PF
        for s in entity.get("actuals", []):
            if _is_grand_total_section(s): continue
            for item in s.get("items", []):
                if item.get("name") == iname:
                    return float(item.get(f, 0) or 0)
            for sub in s.get("sub_heads", []):
                for item in sub.get("items", []):
                    if item.get("name") == iname:
                        return float(item.get(f, 0) or 0)
        return 0.0

    def _grand_v(entity, f):
        # Sum all non-GRAND-TOTAL sections to avoid double-counting
        v = 0.0
        for s in entity.get("actuals", []):
            if _is_grand_total_section(s): continue
            v += float(s.get(f, 0) or 0)
        return v

    def _write_row(head_val, type_val, vp_list, ve_list, fill, font):
        ws.append(["", head_val, type_val] + [""] * (last_col - c_type))
        r = ws.max_row
        gp = ge = 0.0
        for ei in range(n):
            cp = ws.cell(r, c_ent_start + ei*2,     vp_list[ei] if vp_list[ei] else None)
            ce = ws.cell(r, c_ent_start + ei*2 + 1, ve_list[ei] if ve_list[ei] else None)
            cp.number_format = NUM_FMT_CR; ce.number_format = NUM_FMT_CR
            gp += vp_list[ei]; ge += ve_list[ei]
        ws.cell(r, c_grand_p, gp if gp else None).number_format = NUM_FMT_CR
        ws.cell(r, c_grand_e, ge if ge else None).number_format = NUM_FMT_CR
        for col in range(2, last_col+1):
            c = ws.cell(r, col)
            c.fill = fill; c.font = font; c.border = brd
            c.alignment = _LEFT if col <= c_type else _RIGHT

    # ── Sheet title ──
    ws.append(["", org_name])
    _merge(ws, 1, 2, 1, last_col)
    ws["B1"].font = Font(size=14, bold=True, name="Calibri", color="003B63")
    ws["B1"].alignment = _LEFT
    ws.append(["", f"Budget & Estimate – {fy}"])
    _merge(ws, 2, 2, 2, last_col)
    ws["B2"].font = Font(size=12, bold=True, name="Calibri"); ws["B2"].alignment = _LEFT
    ws.append([])

    # Note
    ws.cell(4, last_col, "₹ (Absolute)")
    ws.cell(4, last_col).font = _ITALIC; ws.cell(4, last_col).alignment = _RIGHT
    _merge(ws, 4, 2, 4, last_col)

    # ── Headers — style ALL cells BEFORE merging ──
    for col in range(2, last_col+1):
        for hr in [5, 6]:
            c = ws.cell(hr, col); c.fill = _FILL_HDR; c.font = _WHITE_BOLD
            c.border = brd; c.alignment = _CENTER

    ws.cell(5, c_head, "HEAD OF EXPENSE")
    ws.cell(5, c_type, "TYPE OF EXPENSE")
    _merge(ws, 5, c_head, 6, c_head)
    _merge(ws, 5, c_type, 6, c_type)

    for ei, entity in enumerate(entities):
        cs = c_ent_start + ei*2
        ws.cell(5, cs, (entity.get("label") or "").strip())
        _merge(ws, 5, cs, 5, cs+1)
        ws.cell(6, cs,   plan_label)
        ws.cell(6, cs+1, est_label)

    ws.cell(5, c_grand_p, "GRAND TOTAL")
    ws.cell(5, c_grand_p).fill = PatternFill("solid", fgColor="003B63")
    ws.cell(5, c_grand_e).fill = PatternFill("solid", fgColor="003B63")
    _merge(ws, 5, c_grand_p, 5, c_grand_e)
    ws.cell(6, c_grand_p, plan_label)
    ws.cell(6, c_grand_e, est_label)
    ws.cell(6, c_grand_p).fill = PatternFill("solid", fgColor="004F8B")
    ws.cell(6, c_grand_e).fill = PatternFill("solid", fgColor="004F8B")

    ws.freeze_panes = ws.cell(7, c_type+1).coordinate

    if not entities:
        _auto_col_widths(ws); return ws

    for sec in entities[0].get("actuals", []):
        if _is_grand_total_section(sec): continue
        sname = sec.get("name", "")

        # Section header
        ws.append(["", sname])
        r_sec = ws.max_row
        for col in range(2, last_col+1):
            c = ws.cell(r_sec, col)
            c.fill = _FILL_HEAD
            c.font = Font(bold=True, name="Calibri", size=10, color="003B63")
            c.border = brd; c.alignment = _LEFT
        _merge(ws, r_sec, 2, r_sec, last_col)

        for sub in sec.get("sub_heads", []):
            subname = sub.get("name","")
            vp = [_sub_v(e, sname, subname, PF)  for e in entities]
            ve = [_sub_v(e, sname, subname, EF)  for e in entities]
            _write_row(sname, subname, vp, ve, _FILL_SUBHEAD,
                       Font(bold=True, name="Calibri", size=10, color="7A3B00"))
            for item in sub.get("items", []):
                iname = item.get("name","")
                vp2 = [_item_v(e, iname, False) for e in entities]
                ve2 = [_item_v(e, iname, True)  for e in entities]
                _write_row("", iname, vp2, ve2, _FILL_WHITE, _NORMAL)

        for item in sec.get("items", []):
            iname = item.get("name","")
            vp = [_item_v(e, iname, False) for e in entities]
            ve = [_item_v(e, iname, True)  for e in entities]
            _write_row(sname, iname, vp, ve, _FILL_WHITE, _NORMAL)

        # Section total
        vp = [_sec_v(e, sname, PF) for e in entities]
        ve = [_sec_v(e, sname, EF) for e in entities]
        _write_row(f"TOTAL – {sname}", "", vp, ve, _FILL_HTOTAL,
                   Font(bold=True, name="Calibri", size=10, color="003B63"))

    # Grand total
    vp = [_grand_v(e, PF) for e in entities]
    ve = [_grand_v(e, EF) for e in entities]
    _write_row("GRAND TOTAL", "", vp, ve, _FILL_GRAND,
               Font(bold=True, color="FFFFFF", name="Calibri", size=10))

    _auto_col_widths(ws)
    return ws


# # ═══════════════════════════════════════════════════════════════════════════════
# # FY HELPERS
# # ═══════════════════════════════════════════════════════════════════════════════

# def _fy_labels(fy):
#     parts    = (fy or "2025-26").split("-")
#     start_yy = (parts[0] or "2025")[-2:]
#     end_yy   = (parts[1] if len(parts) > 1 else "26")[-2:]
#     prev_s   = str(int(start_yy) - 1).zfill(2)
#     prev_e   = str(int(end_yy)   - 1).zfill(2)
#     return {
#         "plan": f"FY{start_yy}-{end_yy} Plan",
#         "est":  f"FY{prev_s}-{prev_e} Estimate",
#     }

# def _prev_fy(fy):
#     parts = (fy or "2025-26").split("-")
#     s = int(parts[0] or 2025) - 1
#     e = int(parts[1] or 26)   - 1
#     return f"{s}-{str(e).zfill(2)}"


# # ═══════════════════════════════════════════════════════════════════════════════
# # WHITELISTED API ENDPOINTS
# # ═══════════════════════════════════════════════════════════════════════════════

# @frappe.whitelist()
# def export_ppt(financial_year, ppt_rows, prev_ppt_rows,
#                budget_label, est_label, prev_budget_label, prev_est_label):
#     fy      = financial_year or "2025-26"
#     prev_fy = _prev_fy(fy)
#     wb = Workbook(); wb.remove(wb.active)
#     _sheet_ppt_combined(
#         wb, "Foundation Metrics",
#         json.loads(ppt_rows),      budget_label,      est_label,
#         json.loads(prev_ppt_rows), prev_budget_label, prev_est_label,
#         fy=fy, prev_fy=prev_fy,
#     )
#     return {"filename": f"Foundation_Metrics_{fy}.xlsx", "data": _wb_to_b64(wb)}


# @frappe.whitelist()
# def export_summary_inr(financial_year, summary_data):
#     fy = financial_year or "2025-26"
#     wb = Workbook(); wb.remove(wb.active)
#     _sheet_summary_inr(wb, "Summary in INR", json.loads(summary_data), fy)
#     return {"filename": f"Summary_INR_{fy}.xlsx", "data": _wb_to_b64(wb)}


# @frappe.whitelist()
# def export_headcount(financial_year, headcount_data):
#     fy = financial_year or "2025-26"
#     wb = Workbook(); wb.remove(wb.active)
#     _sheet_headcount(wb, "Headcount", json.loads(headcount_data), fy)
#     return {"filename": f"Headcount_{fy}.xlsx", "data": _wb_to_b64(wb)}


# @frappe.whitelist()
# def export_annual(financial_year, annual_data):
#     fy = financial_year or "2025-26"
#     wb = Workbook(); wb.remove(wb.active)
#     _sheet_annual(wb, "Annual Budget Consolidated", json.loads(annual_data), fy)
#     return {"filename": f"Annual_Budget_Consolidated_{fy}.xlsx", "data": _wb_to_b64(wb)}


# @frappe.whitelist()
# def export_estimate(financial_year, estimate_data):
#     fy = financial_year or "2025-26"
#     wb = Workbook(); wb.remove(wb.active)
#     _sheet_estimate(wb, "Estimate Consolidated", json.loads(estimate_data), fy)
#     return {"filename": f"Estimate_Consolidated_{fy}.xlsx", "data": _wb_to_b64(wb)}


# @frappe.whitelist()
# def export_budget_estimate(financial_year, be_data):
#     fy     = financial_year or "2025-26"
#     labels = _fy_labels(fy)
#     wb = Workbook(); wb.remove(wb.active)
#     _sheet_be(wb, "Budget & Estimate", json.loads(be_data),
#               fy, labels["plan"], labels["est"])
#     return {"filename": f"Budget_and_Estimate_{fy}.xlsx", "data": _wb_to_b64(wb)}


# @frappe.whitelist()
# def export_all(financial_year,
#                ppt_rows, prev_ppt_rows,
#                budget_label, est_label, prev_budget_label, prev_est_label,
#                summary_data, headcount_data,
#                annual_data, estimate_data, be_data):
#     fy      = financial_year or "2025-26"
#     prev_fy = _prev_fy(fy)
#     labels  = _fy_labels(fy)
#     wb = Workbook(); wb.remove(wb.active)

#     _sheet_ppt_combined(
#         wb, "Foundation Metrics",
#         json.loads(ppt_rows),      budget_label,      est_label,
#         json.loads(prev_ppt_rows), prev_budget_label, prev_est_label,
#         fy=fy, prev_fy=prev_fy,
#     )
#     _sheet_summary_inr(wb, "Summary in INR",            json.loads(summary_data),    fy)
#     _sheet_headcount(  wb, "Headcount",                 json.loads(headcount_data),  fy)
#     _sheet_annual(     wb, "Annual Budget Consolidated", json.loads(annual_data),     fy)
#     _sheet_estimate(   wb, "Estimate Consolidated",      json.loads(estimate_data),   fy)
#     _sheet_be(         wb, "Budget & Estimate",          json.loads(be_data),
#               fy, labels["plan"], labels["est"])

#     return {
#         "filename": f"Foundation_Consolidated_Budget_{fy}.xlsx",
#         "data": _wb_to_b64(wb)
#     }



import io
import base64
import json

import frappe
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ═══════════════════════════════════════════════════════════════════════════════
# SHARED STYLE CONSTANTS  — colours match JS CSS variables exactly
# ═══════════════════════════════════════════════════════════════════════════════

_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=False)
_LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
_RIGHT  = Alignment(horizontal="right",  vertical="center")

_BOLD       = Font(bold=True,  name="Calibri", size=10)
_NORMAL     = Font(bold=False, name="Calibri", size=10)
_WHITE_BOLD = Font(bold=True,  color="FFFFFF", name="Calibri", size=10)
_ITALIC     = Font(bold=False, italic=True,    name="Calibri", size=9, color="555555")

# Exact colours from JS --blue-mid, --orange, --blue-light, --orange-light, --blue-dark
_FILL_BLUE_MID    = PatternFill("solid", fgColor="0076B6")   # main header
_FILL_ORANGE      = PatternFill("solid", fgColor="F26B21")   # sub-header
_FILL_BLUE_LIGHT  = PatternFill("solid", fgColor="E9F4FB")   # cb-row-head
_FILL_ORANGE_LIGHT= PatternFill("solid", fgColor="FFF3E6")   # cb-row-sub
_FILL_BLUE_DARK   = PatternFill("solid", fgColor="003B63")   # cb-row-grand
_FILL_HTOTAL      = PatternFill("solid", fgColor="E8F0FA")   # ppt-total-row / sinr-total-row
_FILL_GT          = PatternFill("solid", fgColor="DDEAF7")   # sinr-gt-plan / sinr-gt-act
_FILL_WHITE       = PatternFill("solid", fgColor="FFFFFF")
_FILL_FAFAFA      = PatternFill("solid", fgColor="FAFAFA")   # sinr-brkdwn-act
_FILL_SPACER      = PatternFill("solid", fgColor="F4F6F8")   # sinr-spacer

_THIN   = Side(style="thin")
_MEDIUM = Side(style="medium")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_THICK  = Border(left=_THIN, right=_THIN, top=_MEDIUM, bottom=_THIN)

# Number formats — JS fmtCr divides by 1e7 and shows 2dp
NUM_FMT_CR  = '#,##0.00'   # Crore values
NUM_FMT_INT = '#,##0'      # Headcount whole numbers
NUM_FMT_PCT = '0.0%'       # percentage (stored as decimal, e.g. 0.357)
NUM_FMT_RAW = '#,##0.00'   # Annual/Estimate raw rupees


# ═══════════════════════════════════════════════════════════════════════════════
# LOW-LEVEL HELPERS
# ═══════════════════════════════════════════════════════════════════════════════

def _to_cr(v):
    """Raw rupees → Crores (÷ 1,00,00,000), matching JS fmtCr."""
    return round(float(v or 0) / 1e7, 2)


def _fv(v):
    return float(v or 0)


def _norm(s):
    return (s or "").replace("  ", " ").strip().upper()


def _is_gt(sec):
    nm = _norm(sec.get("name", ""))
    return sec.get("sequence_id") == 9999 or nm == "GRAND TOTAL"


def _is_consolidated(entry):
    return (entry.get("sequence_id") == 9999
            or (entry.get("table_name") or "").upper() == "CONSOLIDATED")


def _wb_to_b64(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return base64.b64encode(buf.read()).decode("utf-8")


def _auto_widths(ws, min_w=8, max_w=55):
    for col in range(1, ws.max_column + 1):
        best = min_w
        for row in range(1, ws.max_row + 1):
            v = ws.cell(row, col).value
            if v is not None:
                s = str(v)
                if s.startswith("="):
                    s = "999,999,999.00"
                best = max(best, len(s))
        ws.column_dimensions[get_column_letter(col)].width = min(best + 2, max_w)


def _cell(ws, row, col, value=None, fill=None, font=None, align=None, border=None, fmt=None):
    c = ws.cell(row=row, column=col, value=value)
    if fill:   c.fill   = fill
    if font:   c.font   = font
    if align:  c.alignment = align
    if border: c.border = border
    if fmt:    c.number_format = fmt
    return c


def _row_style(ws, row, cols, fill=None, font=None, align=_RIGHT, border=_BORDER):
    for col in cols:
        c = ws.cell(row, col)
        if fill:  c.fill  = fill
        if font:  c.font  = font
        if align: c.alignment = align
        c.border = border


def _merge(ws, r1, c1, r2, c2):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def _fy_labels(fy):
    p  = (fy or "2025-26").split("-")
    sY = (p[0] or "2025")[-2:]
    eY = (p[1] if len(p) > 1 else "26")[-2:]
    ps = str(int(sY) - 1).zfill(2)
    pe = str(int(eY) - 1).zfill(2)
    return {"plan": f"FY{sY}-{eY} Plan", "est": f"FY{ps}-{pe} Estimate"}


def _prev_fy(fy):
    p = (fy or "2025-26").split("-")
    s = int(p[0] or 2025) - 1
    e = int(p[1] or 26) - 1
    return f"{s}-{str(e).zfill(2)}"


def _sheet_title(ws, org, subtitle, last_col):
    """Write 2 title rows at top of sheet."""
    c = ws.cell(1, 2, org)
    c.font = Font(size=13, bold=True, name="Calibri", color="003B63")
    c.alignment = _LEFT
    _merge(ws, 1, 2, 1, last_col)
    c2 = ws.cell(2, 2, subtitle)
    c2.font = Font(size=11, bold=True, name="Calibri")
    c2.alignment = _LEFT
    _merge(ws, 2, 2, 2, last_col)
    ws.append([])


# ═══════════════════════════════════════════════════════════════════════════════
# SHEET 1 — Foundation Level / Overall Metrics  (PPT tab)
#
# JS display:
#   Two tables stacked: current FY and previous FY
#   Columns: Unit | Budget(Opex Capex Total) | Actual(Opex Capex Total)
#   Total uses GRAND TOTAL section (seq 9999) directly for the Total column
#   Opex / Capex come from OPERATING/CAPITAL named sections
#   Last row = "Total" with blue background
# ═══════════════════════════════════════════════════════════════════════════════

def _extract_opex_capex(sections, field):
    """
    Match JS extractVals: sum OPERATING→opex, CAPITAL→capex.
    Fallback to GRAND TOTAL section if no breakdown found.
    """
    opex = capex = 0.0
    has_breakdown = False
    for sec in (sections or []):
        nm = _norm(sec.get("name", ""))
        if sec.get("sequence_id") == 9999 or nm == "GRAND TOTAL":
            continue
        if "OPERATING" in nm:
            opex += _fv(sec.get(field)); has_breakdown = True
        elif "CAPITAL" in nm:
            capex += _fv(sec.get(field)); has_breakdown = True
    if not has_breakdown:
        for sec in (sections or []):
            nm = _norm(sec.get("name", ""))
            if sec.get("sequence_id") == 9999 or nm == "GRAND TOTAL":
                opex += _fv(sec.get(field))
        if not opex:
            for sec in (sections or []):
                nm = _norm(sec.get("name", ""))
                if sec.get("sequence_id") != 9999 and nm != "GRAND TOTAL":
                    opex += _fv(sec.get(field))
    return opex, capex


def _extract_total(sections, field):
    """
    Match JS extractTotal: use GRAND TOTAL section (seq 9999) directly.
    Fallback: sum all sections.
    """
    gt = 0.0
    found = False
    for sec in (sections or []):
        nm = _norm(sec.get("name", ""))
        if sec.get("sequence_id") == 9999 or nm == "GRAND TOTAL":
            gt += _fv(sec.get(field)); found = True
    if not found:
        for sec in (sections or []):
            gt += _fv(sec.get(field))
    return gt


def _sheet_ppt(wb, ppt_rows, budget_label, est_label,
               prev_rows, prev_budget_label, prev_est_label, fy, prev_fy,
               org="Azim Premji Foundation"):
    ws = wb.create_sheet("Foundation Metrics")
    COLS = list(range(2, 9))   # B–H  (7 data cols)

    _sheet_title(ws, org, f"Foundation Level / Overall Metrics – {fy}", 8)

    def _write_block(start_row, rows, b_lbl, e_lbl, title):
        # Title bar
        ws.cell(start_row, 2, title)
        for col in COLS:
            c = ws.cell(start_row, col)
            c.fill = _FILL_BLUE_LIGHT
            c.font = Font(bold=True, name="Calibri", size=11, color="003B63")
            c.border = _BORDER
            c.alignment = _LEFT
        _merge(ws, start_row, 2, start_row, 8)

        # Currency note
        nr = start_row + 1
        ws.cell(nr, 8, "₹ Cr.").font = _ITALIC
        ws.cell(nr, 8).alignment = Alignment(horizontal="right")

        # Header row 1 (blue)
        h1 = start_row + 2
        for col in COLS:
            c = ws.cell(h1, col)
            c.fill = _FILL_BLUE_MID; c.font = _WHITE_BOLD
            c.border = _BORDER; c.alignment = _CENTER
        ws.cell(h1, 2, "Unit")
        ws.cell(h1, 3, b_lbl)
        ws.cell(h1, 6, e_lbl)
        _merge(ws, h1, 2, h1 + 1, 2)
        _merge(ws, h1, 3, h1, 5)
        _merge(ws, h1, 6, h1, 8)

        # Header row 2 (orange)
        h2 = start_row + 3
        for col in COLS:
            c = ws.cell(h2, col)
            c.fill = _FILL_ORANGE; c.font = _WHITE_BOLD
            c.border = _BORDER; c.alignment = _CENTER
        for col, lbl in zip(range(3, 9), ["Opex", "Capex", "Total", "Opex", "Capex", "Total"]):
            ws.cell(h2, col, lbl)

        # Data rows
        dr = start_row + 4
        for row in rows:
            is_total = row.get("is_total", False)
            fill = _FILL_HTOTAL if is_total else _FILL_WHITE
            font = _BOLD if is_total else _NORMAL
            bO = _to_cr(row.get("bOpex",  0))
            bC = _to_cr(row.get("bCapex", 0))
            bT = _to_cr(row.get("bTotal", 0)) if row.get("bTotal") else _to_cr(_fv(row.get("bOpex", 0)) + _fv(row.get("bCapex", 0)))
            eO = _to_cr(row.get("eOpex",  0))
            eC = _to_cr(row.get("eCapex", 0))
            eT = _to_cr(row.get("eTotal", 0)) if row.get("eTotal") else _to_cr(_fv(row.get("eOpex", 0)) + _fv(row.get("eCapex", 0)))
            for col, val in enumerate([row.get("label", ""), bO or None, bC or None, bT or None,
                                        eO or None, eC or None, eT or None], start=2):
                c = ws.cell(dr, col, val)
                c.fill = fill; c.font = font; c.border = _BORDER
                c.alignment = _LEFT if col == 2 else _RIGHT
                if col > 2:
                    c.number_format = NUM_FMT_CR
            dr += 1
        return dr

    next_r = _write_block(4, ppt_rows, budget_label, est_label,
                          f"Overall Foundation – {fy} Budget vs {prev_fy} Actual")
    _write_block(next_r + 1, prev_rows, prev_budget_label, prev_est_label,
                 f"Overall Foundation – {prev_fy} Budget vs {prev_fy} Actual")

    _auto_widths(ws)
    return ws


# ═══════════════════════════════════════════════════════════════════════════════
# SHEET 2 — Summary in INR  (SummaryINR tab)
#
# JS display:
#   Table A: Unit / Function | Plan(Opex Capex Total) | Act(Opex Capex Total)
#     - Each unit row from API
#     - "Grand Total" row from CONSOLIDATED TOTAL entry (seq 9999)
#     - Values in Crores (÷ 1e7)
#
#   Table B: Breakdown — rows per unit, plan + actual rows
#     - Columns: Unit | Opex sub-heads… | Opex Total | Capex | Total
#     - Grand Total uses CONSOLIDATED TOTAL values
# ═══════════════════════════════════════════════════════════════════════════════

def _sinr_extract_a(actuals):
    """Match JS extractA: sum OPERATING→opex_plan/act, CAPITAL→capex_plan/act."""
    op_p = cp_p = op_a = cp_a = 0.0
    for sec in (actuals or []):
        nm = _norm(sec.get("name", ""))
        if "OPERATING" in nm:
            op_p += _fv(sec.get("ytd"))
            op_a += _fv(sec.get("total_posted_amt_ytd"))
        elif "CAPITAL" in nm:
            cp_p += _fv(sec.get("ytd"))
            cp_a += _fv(sec.get("total_posted_amt_ytd"))
    return op_p, cp_p, op_a, cp_a


def _sinr_get_consolidated(data):
    """
    Match JS getConsolidatedTotals: find seq_id 9999 / table CONSOLIDATED entry.
    Returns dict with opex_plan, capex_plan, total_plan, opex_act, capex_act, total_act.
    """
    for e in (data or []):
        if _is_consolidated(e):
            r = {"opex_plan": 0, "capex_plan": 0, "total_plan": 0,
                 "opex_act":  0, "capex_act":  0, "total_act":  0}
            for a in (e.get("actuals") or []):
                nm = _norm(a.get("name", ""))
                if nm == "OPEX TOTAL":
                    r["opex_plan"] += _fv(a.get("ytd"))
                    r["opex_act"]  += _fv(a.get("total_posted_amt_ytd"))
                elif nm == "CAPEX TOTAL":
                    r["capex_plan"] += _fv(a.get("ytd"))
                    r["capex_act"]  += _fv(a.get("total_posted_amt_ytd"))
                elif nm == "OVERALL GRAND TOTAL":
                    r["total_plan"] = _fv(a.get("ytd"))
                    r["total_act"]  = _fv(a.get("total_posted_amt_ytd"))
            if not r["total_plan"] and not r["total_act"]:
                r["total_plan"] = r["opex_plan"] + r["capex_plan"]
                r["total_act"]  = r["opex_act"]  + r["capex_act"]
            return r
    return None


def _sheet_summary_inr(wb, api_data, fy, org="Azim Premji Foundation"):
    ws = wb.create_sheet("Summary in INR")
    fp = (fy or "2025-26").split("-")
    p_lbl = f"{fy} Budget"
    a_lbl = f"{int(fp[0])-1}-{str(int(fp[1])-1).zfill(2)} Est"

    _sheet_title(ws, org, f"Summary in INR – {fy}", 8)

    sorted_data = sorted((api_data or []), key=lambda e: e.get("sequence_id", 0))
    # Exclude consolidated total entry from unit rows
    unit_rows  = [e for e in sorted_data if not _is_consolidated(e)]
    normal     = [e for e in unit_rows if "covid" not in (e.get("label") or "").lower()]
    covid_rows = [e for e in unit_rows if "covid"     in (e.get("label") or "").lower()]
    ct         = _sinr_get_consolidated(api_data)

    cur = 4  # current row

    # ── Table A header ──────────────────────────────────────────────────────
    ws.cell(cur, 2, "A. Unit Wise Plan")
    ws.cell(cur, 2).font = Font(bold=True, name="Calibri", size=11, underline="single")
    cur += 1
    ws.cell(cur, 8, "₹ Cr."); ws.cell(cur, 8).font = _ITALIC
    ws.cell(cur, 8).alignment = Alignment(horizontal="right"); cur += 1

    h1 = cur
    for col in range(2, 9):
        c = ws.cell(h1, col)
        c.fill = _FILL_BLUE_MID; c.font = _WHITE_BOLD; c.border = _BORDER; c.alignment = _CENTER
    ws.cell(h1, 2, "Unit / Function"); ws.cell(h1, 3, p_lbl); ws.cell(h1, 6, a_lbl)
    _merge(ws, h1, 2, h1 + 1, 2); _merge(ws, h1, 3, h1, 5); _merge(ws, h1, 6, h1, 8)
    cur += 1

    h2 = cur
    for col in range(2, 9):
        c = ws.cell(h2, col)
        c.fill = _FILL_ORANGE; c.font = _WHITE_BOLD; c.border = _BORDER; c.alignment = _CENTER
    for col, lbl in zip(range(3, 9), ["Opex", "Capex", "Total", "Opex", "Capex", "Total"]):
        ws.cell(h2, col, lbl)
    cur += 1

    def _write_unit_row_a(entry, fill=_FILL_WHITE, font=_NORMAL, label_override=None):
        nonlocal cur
        op_p, cp_p, op_a, cp_a = _sinr_extract_a(entry.get("actuals", []))
        is_sub = entry.get("is_this_sub_item") == 1
        lbl    = label_override or (entry.get("label") or "").strip()
        indent = "    " if is_sub else ""
        for col, val in enumerate([indent + lbl,
                                    _to_cr(op_p), _to_cr(cp_p), _to_cr(op_p + cp_p),
                                    _to_cr(op_a), _to_cr(cp_a), _to_cr(op_a + cp_a)], start=2):
            c = ws.cell(cur, col, val if (col == 2 or val != 0.0) else None)
            c.fill = fill; c.font = font; c.border = _BORDER
            c.alignment = _LEFT if col == 2 else _RIGHT
            if col > 2: c.number_format = NUM_FMT_CR
        cur += 1

    for e in normal:
        _write_unit_row_a(e)

    for e in covid_rows:
        _write_unit_row_a(e)

    # Grand Total row (from CONSOLIDATED TOTAL entry)
    if ct:
        op_p = ct["opex_plan"]; cp_p = ct["capex_plan"]; tp = ct["total_plan"]
        op_a = ct["opex_act"];  cp_a = ct["capex_act"];  ta = ct["total_act"]
    else:
        op_p = cp_p = op_a = cp_a = 0.0
        for e in unit_rows:
            if e.get("is_this_sub_item") != 1:
                o, c2, oa, ca = _sinr_extract_a(e.get("actuals", []))
                op_p += o; cp_p += c2; op_a += oa; cp_a += ca
        tp = op_p + cp_p; ta = op_a + cp_a

    for col, val in enumerate(["Grand Total",
                                _to_cr(op_p), _to_cr(cp_p), _to_cr(tp),
                                _to_cr(op_a), _to_cr(cp_a), _to_cr(ta)], start=2):
        c = ws.cell(cur, col, val)
        c.fill = _FILL_BLUE_DARK
        c.font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
        c.border = _BORDER
        c.alignment = _LEFT if col == 2 else _RIGHT
        if col > 2: c.number_format = NUM_FMT_CR
    cur += 2

    # ── Table B header ──────────────────────────────────────────────────────
    ws.cell(cur, 2, "B. Breakdown of Unit Wise Plan")
    ws.cell(cur, 2).font = Font(bold=True, name="Calibri", size=11, underline="single")
    cur += 1
    ws.cell(cur, 8, "₹ Cr."); ws.cell(cur, 8).font = _ITALIC
    ws.cell(cur, 8).alignment = Alignment(horizontal="right"); cur += 1

    # Collect sub-head names from OPERATING section (match JS getSubNames)
    main_entries = [e for e in unit_rows if e.get("is_this_sub_item") != 1]
    seen_sh = {}; sub_head_names = []
    for e in main_entries:
        for sec in (e.get("actuals") or []):
            if "OPERATING" not in _norm(sec.get("name", "")): continue
            for sh in (sec.get("sub_heads") or []):
                n = (sh.get("name") or "").strip()
                if n and n not in seen_sh: seen_sh[n] = True; sub_head_names.append(n)

    n_sh    = len(sub_head_names)
    c_start = 3                     # first sub-head col
    c_opx_t = c_start + n_sh        # Opex Total col
    c_capex = c_opx_t + 1           # Capex col
    c_total = c_capex + 1           # Total col
    last_c  = c_total

    # Header row 1
    bh1 = cur
    for col in range(2, last_c + 1):
        c = ws.cell(bh1, col)
        c.fill = _FILL_BLUE_MID; c.font = _WHITE_BOLD; c.border = _BORDER; c.alignment = _CENTER
    ws.cell(bh1, 2, "Unit / Function"); _merge(ws, bh1, 2, bh1 + 1, 2)
    ws.cell(bh1, c_start, "Operating Expenses")
    _merge(ws, bh1, c_start, bh1, c_opx_t)
    ws.cell(bh1, c_capex, "Capex"); _merge(ws, bh1, c_capex, bh1 + 1, c_capex)
    ws.cell(bh1, c_total, "Total"); _merge(ws, bh1, c_total, bh1 + 1, c_total)
    cur += 1

    # Header row 2
    bh2 = cur
    for col in range(2, last_c + 1):
        c = ws.cell(bh2, col)
        c.fill = _FILL_ORANGE; c.font = _WHITE_BOLD; c.border = _BORDER; c.alignment = _CENTER
    for i, shn in enumerate(sub_head_names):
        ws.cell(bh2, c_start + i, shn)
    ws.cell(bh2, c_opx_t, "Total")
    cur += 1

    def _sh_val(actuals, shname, field):
        v = 0.0
        for sec in (actuals or []):
            if "OPERATING" not in _norm(sec.get("name", "")): continue
            for sh in (sec.get("sub_heads") or []):
                if sh.get("name", "").strip() == shname: v += _fv(sh.get(field))
        return v

    def _opex_total(actuals, field):
        v = 0.0
        for sec in (actuals or []):
            if "OPERATING" not in _norm(sec.get("name", "")): continue
            v += _fv(sec.get(field))
        return v

    def _capex_total(actuals, field):
        v = 0.0
        for sec in (actuals or []):
            if "CAPITAL" not in _norm(sec.get("name", "")): continue
            v += _fv(sec.get(field))
        return v

    gt_sh_p = {n: 0.0 for n in sub_head_names}
    gt_sh_a = {n: 0.0 for n in sub_head_names}
    gt_op_p = gt_op_a = gt_cp_p = gt_cp_a = 0.0

    for entry in main_entries:
        act = entry.get("actuals") or []
        lbl = (entry.get("label") or "").strip()

        # Unit header row — sinr-unit-hdr (blue light)
        for col in range(2, last_c + 1):
            c = ws.cell(cur, col)
            c.fill = _FILL_BLUE_LIGHT; c.font = _BOLD; c.border = _BORDER
            c.alignment = _LEFT if col == 2 else _CENTER
        ws.cell(cur, 2, lbl); _merge(ws, cur, 3, cur, last_c)
        cur += 1

        # Plan row — sinr-brkdwn-plan
        op_p = _opex_total(act, "ytd"); cp_p = _capex_total(act, "ytd")
        for col in range(2, last_c + 1):
            c = ws.cell(cur, col)
            c.fill = _FILL_WHITE; c.font = _NORMAL; c.border = _BORDER
            c.alignment = _LEFT if col == 2 else _RIGHT
        ws.cell(cur, 2, f"  - {p_lbl}")
        for i, shn in enumerate(sub_head_names):
            v = _sh_val(act, shn, "ytd")
            gt_sh_p[shn] += v
            c = ws.cell(cur, c_start + i, _to_cr(v) or None)
            c.number_format = NUM_FMT_CR
        ws.cell(cur, c_opx_t, _to_cr(op_p) or None).number_format = NUM_FMT_CR
        ws.cell(cur, c_capex, _to_cr(cp_p) or None).number_format = NUM_FMT_CR
        ws.cell(cur, c_total, _to_cr(op_p + cp_p) or None).number_format = NUM_FMT_CR
        gt_op_p += op_p; gt_cp_p += cp_p
        cur += 1

        # Actual row — sinr-brkdwn-act
        op_a = _opex_total(act, "total_posted_amt_ytd"); cp_a = _capex_total(act, "total_posted_amt_ytd")
        for col in range(2, last_c + 1):
            c = ws.cell(cur, col)
            c.fill = _FILL_FAFAFA; c.font = _NORMAL; c.border = _BORDER
            c.alignment = _LEFT if col == 2 else _RIGHT
        ws.cell(cur, 2, f"  - {a_lbl}")
        for i, shn in enumerate(sub_head_names):
            v = _sh_val(act, shn, "total_posted_amt_ytd")
            gt_sh_a[shn] += v
            c = ws.cell(cur, c_start + i, _to_cr(v) or None)
            c.number_format = NUM_FMT_CR
        ws.cell(cur, c_opx_t, _to_cr(op_a) or None).number_format = NUM_FMT_CR
        ws.cell(cur, c_capex, _to_cr(cp_a) or None).number_format = NUM_FMT_CR
        ws.cell(cur, c_total, _to_cr(op_a + cp_a) or None).number_format = NUM_FMT_CR
        gt_op_a += op_a; gt_cp_a += cp_a
        cur += 1

        # Spacer
        for col in range(2, last_c + 1):
            ws.cell(cur, col).fill = _FILL_SPACER; ws.cell(cur, col).border = _BORDER
        cur += 1

    # Grand Total label row — cb-row-grand (dark blue)
    for col in range(2, last_c + 1):
        c = ws.cell(cur, col)
        c.fill = _FILL_BLUE_DARK
        c.font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
        c.border = _BORDER; c.alignment = _LEFT if col == 2 else _CENTER
    ws.cell(cur, 2, "Grand Total"); _merge(ws, cur, 3, cur, last_c)
    cur += 1

    # Use CONSOLIDATED TOTAL values for GT if available (match JS finalOP etc.)
    fin_op_p = ct["opex_plan"]  if ct else gt_op_p
    fin_op_a = ct["opex_act"]   if ct else gt_op_a
    fin_cp_p = ct["capex_plan"] if ct else gt_cp_p
    fin_cp_a = ct["capex_act"]  if ct else gt_cp_a
    fin_tp   = ct["total_plan"] if ct else (gt_op_p + gt_cp_p)
    fin_ta   = ct["total_act"]  if ct else (gt_op_a + gt_cp_a)

    # GT Plan row — sinr-gt-plan
    for col in range(2, last_c + 1):
        c = ws.cell(cur, col); c.fill = _FILL_GT; c.font = _NORMAL
        c.border = _BORDER; c.alignment = _LEFT if col == 2 else _RIGHT
    ws.cell(cur, 2, f"  - {p_lbl}")
    for i, shn in enumerate(sub_head_names):
        ws.cell(cur, c_start + i, _to_cr(gt_sh_p[shn]) or None).number_format = NUM_FMT_CR
    ws.cell(cur, c_opx_t, _to_cr(fin_op_p) or None).number_format = NUM_FMT_CR
    ws.cell(cur, c_capex, _to_cr(fin_cp_p) or None).number_format = NUM_FMT_CR
    ws.cell(cur, c_total, _to_cr(fin_tp)   or None).number_format = NUM_FMT_CR
    cur += 1

    # GT Actual row — sinr-gt-act
    for col in range(2, last_c + 1):
        c = ws.cell(cur, col); c.fill = _FILL_GT; c.font = _NORMAL
        c.border = _BORDER; c.alignment = _LEFT if col == 2 else _RIGHT
    ws.cell(cur, 2, f"  - {a_lbl}")
    for i, shn in enumerate(sub_head_names):
        ws.cell(cur, c_start + i, _to_cr(gt_sh_a[shn]) or None).number_format = NUM_FMT_CR
    ws.cell(cur, c_opx_t, _to_cr(fin_op_a) or None).number_format = NUM_FMT_CR
    ws.cell(cur, c_capex, _to_cr(fin_cp_a) or None).number_format = NUM_FMT_CR
    ws.cell(cur, c_total, _to_cr(fin_ta)   or None).number_format = NUM_FMT_CR

    _auto_widths(ws)
    return ws


# ═══════════════════════════════════════════════════════════════════════════════
# SHEET 3 — Headcount
#
# JS display — 5 sections (matching renderSummary + 4 tables):
#   1. Headcount Summary  — Unit | Avg H/C y[-2] | Avg H/C y[-1] | % | Est Opex | Plan Opex | %
#   2. Closing H/C        — Unit | [hc per year straight from API]
#   3. Average H/C        — Unit | [(prev+curr)/2 per year;  i=0 → hc[0]/2]
#   4. Increase Closing % — transitions (curr/prev)-1
#   5. Increase Avg H/C % — transitions using avg values
# ═══════════════════════════════════════════════════════════════════════════════

def _sheet_headcount(wb, payload, fy, org="Azim Premji Foundation"):
    ws = wb.create_sheet("Headcount")

    if isinstance(payload, dict):
        records   = payload.get("headcount_data") or []
        plan_data = payload.get("plan_data") or []
    else:
        records   = payload or []
        plan_data = []

    # Build opex map from plan_data — match JS buildOpexMap
    def _norm_lbl(s): return (s or "").lower().replace("  ", " ").strip()

    opex_map = {}
    for p in (plan_data or []):
        lbl = _norm_lbl(p.get("label", ""))
        op  = None
        for a in (p.get("actuals") or []):
            nm = (a.get("name") or "").strip()
            if nm in ("OPERATING  EXPENSES", "OPERATING EXPENSES"):
                op = a; break
        opex_map[lbl] = {
            "est" : _fv(op.get("total_posted_amt_ytd", 0) if op else 0) / 1e7,
            "plan": _fv(op.get("ytd", 0)                if op else 0) / 1e7,
        }

    # Transform records — match JS transform (straight from API)
    sorted_recs = sorted((records or []), key=lambda r: r.get("financial_year", ""))
    yrs = [r.get("financial_year", "") for r in sorted_recs]

    um = {}
    for rec in sorted_recs:
        for u in (rec.get("units") or []):
            uid = str(u.get("unit", ""))
            if uid not in um:
                um[uid] = {"id": uid, "desc": "", "hc": {}, "seq": int(uid) if uid.isdigit() else 999}
            um[uid]["hc"][rec["financial_year"]] = _fv(u.get("total_headcount") or 0)
            if rec["financial_year"] == yrs[-1]:
                um[uid]["desc"] = (u.get("unit_description") or u.get("description") or "").strip()
    units  = sorted(um.values(), key=lambda u: u["seq"])
    totals = {r["financial_year"]: _fv(r.get("total_head_count") or r.get("total_headcount") or 0)
              for r in sorted_recs}

    # avgHC: i=0 → hc[0]/2; i>0 → (prev+curr)/2 — match JS
    def _avg_hc(hc, i):
        c = hc.get(yrs[i])
        if i == 0:
            return c / 2 if c else None
        p = hc.get(yrs[i - 1])
        return (p + c) / 2 if (p is not None and c is not None) else None

    def _avg_tot(i):
        c = totals.get(yrs[i])
        if i == 0:
            return c / 2 if c else None
        p = totals.get(yrs[i - 1])
        return (p + c) / 2 if (p is not None and c is not None) else None

    def _fym(fy_str):
        p = (fy_str or "").split("-")
        return f"31-Mar-{p[1][-2:] if len(p) > 1 else (p[0] or '')[-2:]}"

    max_col = max(2 + len(yrs), 8)
    _sheet_title(ws, org, f"Headcount – {fy}", max_col)
    cur = 4

    def _section_hdr(title):
        nonlocal cur
        c = ws.cell(cur, 2, title)
        c.font = Font(bold=True, name="Calibri", size=11, underline="single", color="003B63")
        cur += 1

    def _col_hdr(labels):
        nonlocal cur
        for ci, lbl in enumerate(["Unit"] + labels, start=2):
            c = ws.cell(cur, ci, lbl)
            c.fill = _FILL_BLUE_MID; c.font = _WHITE_BOLD
            c.border = _BORDER; c.alignment = _CENTER
        cur += 1

    def _data_row(label, values, fill=_FILL_WHITE, font=_NORMAL, fmt=NUM_FMT_INT, is_grand=False):
        nonlocal cur
        gf = _FILL_BLUE_DARK
        gfont = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
        ws.cell(cur, 2, label).fill = gf if is_grand else fill
        ws.cell(cur, 2).font = gfont if is_grand else font
        ws.cell(cur, 2).border = _BORDER; ws.cell(cur, 2).alignment = _LEFT
        for ci, val in enumerate(values, start=3):
            c = ws.cell(cur, ci)
            c.fill = gf if is_grand else fill
            c.font = gfont if is_grand else font
            c.border = _BORDER; c.alignment = _RIGHT
            if val is not None:
                if fmt == NUM_FMT_PCT:
                    c.value = val / 100   # store as decimal for % format
                    c.number_format = NUM_FMT_PCT
                else:
                    c.value = round(val) if fmt == NUM_FMT_INT else val
                    if fmt: c.number_format = fmt
        cur += 1

    # ── 1. Headcount Summary ────────────────────────────────────────────────
    _section_hdr("Headcount Summary")
    if len(yrs) >= 2:
        i1 = len(yrs) - 2; i2 = len(yrs) - 1
        _col_hdr([yrs[i1], yrs[i2], "%", f"{yrs[i1]} Est", f"{yrs[i2]} Plan", "%"])
        tot_est = tot_plan = 0.0
        for u in units:
            a1 = _avg_hc(u["hc"], i1); a2 = _avg_hc(u["hc"], i2)
            o  = opex_map.get(_norm_lbl(u["desc"]), {"est": 0.0, "plan": 0.0})
            tot_est += o["est"]; tot_plan += o["plan"]
            pct_hc  = round(((a2 / a1) - 1) * 100, 1) if (a1 and a2) else None
            pct_opx = round(((o["plan"] / o["est"]) - 1) * 100, 1) if (o["est"] and o["plan"]) else None
            _data_row(u["desc"], [
                round(a1) if a1 is not None else None,
                round(a2) if a2 is not None else None,
                pct_hc,
                round(o["est"], 2), round(o["plan"], 2), pct_opx,
            ])
        ta1 = _avg_tot(i1); ta2 = _avg_tot(i2)
        pct_t = round(((ta2 / ta1) - 1) * 100, 1) if (ta1 and ta2) else None
        pct_o = round(((tot_plan / tot_est) - 1) * 100, 1) if (tot_est and tot_plan) else None
        _data_row("Total", [
            round(ta1) if ta1 else None, round(ta2) if ta2 else None, pct_t,
            round(tot_est, 2), round(tot_plan, 2), pct_o,
        ], is_grand=True)
    cur += 1

    # ── 2. Closing H/C ──────────────────────────────────────────────────────
    _section_hdr("Closing H/C")
    _col_hdr([_fym(y) for y in yrs])
    for u in units:
        _data_row(u["desc"], [u["hc"].get(y) for y in yrs])
    _data_row("Total", [totals.get(y) for y in yrs], is_grand=True)
    cur += 1

    # ── 3. Average H/C ──────────────────────────────────────────────────────
    _section_hdr("Average H/C")
    _col_hdr([_fym(y) for y in yrs])
    for u in units:
        _data_row(u["desc"], [_avg_hc(u["hc"], i) for i in range(len(yrs))])
    _data_row("Total", [_avg_tot(i) for i in range(len(yrs))], is_grand=True)
    cur += 1

    # ── 4. Increase in Closing H/C (%) ─────────────────────────────────────
    if len(yrs) >= 2:
        _section_hdr("Increase in Closing H/C (%)")
        pairs = [(yrs[i - 1], yrs[i]) for i in range(1, len(yrs))]
        _col_hdr([f"{_fym(yf)} → {_fym(yt)}" for yf, yt in pairs])
        for u in units:
            vals = []
            for yf, yt in pairs:
                a = u["hc"].get(yf); b = u["hc"].get(yt)
                vals.append(round(((b / a) - 1) * 100, 1) if (a and b) else None)
            _data_row(u["desc"], vals, fmt=NUM_FMT_PCT)
        tot_vals = []
        for yf, yt in pairs:
            a = totals.get(yf); b = totals.get(yt)
            tot_vals.append(round(((b / a) - 1) * 100, 1) if (a and b) else None)
        _data_row("Total", tot_vals, fmt=NUM_FMT_PCT, is_grand=True)
        cur += 1

    # ── 5. Increase in Average H/C (%) ─────────────────────────────────────
    if len(yrs) >= 2:
        _section_hdr("Increase in Average H/C (%)")
        pair_idxs = [(i - 1, i) for i in range(1, len(yrs))]
        _col_hdr([f"{yrs[i1]} → {yrs[i2]}" for i1, i2 in pair_idxs])
        for u in units:
            vals = []
            for i1, i2 in pair_idxs:
                prev = _avg_hc(u["hc"], i1); curr = _avg_hc(u["hc"], i2)
                vals.append(round(((curr / prev) - 1) * 100, 1) if (prev and curr) else None)
            _data_row(u["desc"], vals, fmt=NUM_FMT_PCT)
        tot_vals = []
        for i1, i2 in pair_idxs:
            prev = _avg_tot(i1); curr = _avg_tot(i2)
            tot_vals.append(round(((curr / prev) - 1) * 100, 1) if (prev and curr) else None)
        _data_row("Total", tot_vals, fmt=NUM_FMT_PCT, is_grand=True)

    _auto_widths(ws)
    return ws


# ═══════════════════════════════════════════════════════════════════════════════
# SHEET 4 — Annual Budget Consolidated
#
# JS display: Expense Head rows (blue), Sub-head rows (orange), Item rows,
#   Quarter columns (Q1-Q4 monthly + QTR totals + YEAR total)
#   Values in raw rupees (formatINR — whole number Indian grouping)
# ═══════════════════════════════════════════════════════════════════════════════

# Column layout (1-based)
_C_SI    = 2    # Sl #
_C_HEAD  = 3    # HEAD OF EXPENSE
_C_TYPE  = 4    # TYPE OF EXPENSE
_C_START = 5    # Apr (first data col)
_C_END   = 21   # YEAR total
# 5-7=Q1 months, 8-10=Q2, 11-13=Q3, 14-16=Q4, 17=QTR1, 18=QTR2, 19=QTR3, 20=QTR4, 21=YEAR


def _to_roman(n):
    vals = [1000,900,500,400,100,90,50,40,10,9,5,4,1]
    syms = ["M","CM","D","CD","C","XC","L","XL","X","IX","V","IV","I"]
    out = ""; i = 0
    while n > 0:
        for _ in range(n // vals[i]):
            out += syms[i]; n -= vals[i]
        i += 1
    return out


def _qtr_fmls(r):
    return [f"=SUM(E{r}:G{r})", f"=SUM(H{r}:J{r})", f"=SUM(K{r}:M{r})",
            f"=SUM(N{r}:P{r})", f"=SUM(Q{r}:T{r})"]


def _sum_fml(col_letter, rows):
    if not rows: return 0
    return "=" + "+".join(f"{col_letter}{r}" for r in rows)


def _totals_fmls(rows):
    return [_sum_fml(c, rows) for c in list("EFGHIJKLMNOPQRSTU")]


def _style_annual_row(ws, row, fill=None, font=None, is_hdr=False):
    for col in range(_C_SI, _C_END + 1):
        c = ws.cell(row, col)
        c.border = _BORDER
        if is_hdr:
            c.alignment = _CENTER
        else:
            if col == _C_SI:                   c.alignment = _CENTER
            elif col in (_C_HEAD, _C_TYPE):    c.alignment = _LEFT
            else:                              c.alignment = _RIGHT
        if fill: c.fill = fill
        if font: c.font = font


def _fmt_annual_row(ws, row):
    for col in range(_C_START, _C_END + 1):
        ws.cell(row, col).number_format = NUM_FMT_RAW


def _write_annual_headers(ws, fy, r1_label="Annual Budget Consolidated"):
    ws.append([
        "", "Sl #", "HEAD OF EXPENSE", "TYPE OF EXPENSE",
        "QUARTER I", "", "", "QUARTER II", "", "",
        "QUARTER III", "", "", "QUARTER IV", "", "",
        "QTR-1", "QTR-2", "QTR-3", "QTR-4", f"YEAR {fy}",
    ])
    r1 = ws.max_row
    ws.append([
        "", "Sl #", "HEAD OF EXPENSE", "TYPE OF EXPENSE",
        "Apr", "May", "Jun", "Jul", "Aug", "Sep",
        "Oct", "Nov", "Dec", "Jan", "Feb", "Mar",
        "QTR-1", "QTR-2", "QTR-3", "QTR-4", f"YEAR {fy}",
    ])
    r2 = ws.max_row
    _style_annual_row(ws, r1, _FILL_BLUE_MID, _WHITE_BOLD, is_hdr=True)
    _style_annual_row(ws, r2, _FILL_BLUE_MID, _WHITE_BOLD, is_hdr=True)
    for c1, c2 in [(5,7),(8,10),(11,13),(14,16)]:
        _merge(ws, r1, c1, r1, c2)
    for col in [2, 3, 4]:
        _merge(ws, r1, col, r2, col)
    for col in range(17, 22):
        _merge(ws, r1, col, r2, col)
    ws.freeze_panes = "E7"


def _sheet_annual(wb, data, fy, org="Azim Premji Foundation"):
    ws = wb.create_sheet("Annual Budget Consolidated")
    _sheet_title(ws, org, f"Annual Budget Consolidated – {fy}", 21)
    _write_annual_headers(ws, fy)

    head_total_rows = []
    head_counter    = 0

    for head in (data or []):
        head_counter += 1
        alpha = chr(64 + head_counter)
        head_name = _norm(head.get("name") or "")

        if "COVID" in head_name:
            ws.append([])
            item = (head.get("items") or [{}])[0]
            r = ws.max_row + 1
            ws.append(["", alpha, head.get("name", ""), item.get("name", ""),
                        *item.get("q1", [0, 0, 0]), *item.get("q2", [0, 0, 0]),
                        *item.get("q3", [0, 0, 0]), *item.get("q4", [0, 0, 0]),
                        *_qtr_fmls(r)])
            _style_annual_row(ws, ws.max_row, font=_NORMAL)
            _fmt_annual_row(ws, ws.max_row)
            continue

        # Section header row
        ws.append(["", alpha, head.get("name", "")])
        r_sec = ws.max_row
        for col in range(_C_HEAD, _C_END + 1):
            c = ws.cell(r_sec, col)
            c.fill = _FILL_BLUE_LIGHT
            c.font = Font(bold=True, name="Calibri", size=10, color="003B63")
            c.border = _BORDER; c.alignment = _LEFT
        ws.cell(r_sec, _C_SI).fill = _FILL_BLUE_LIGHT
        ws.cell(r_sec, _C_SI).border = _BORDER
        _merge(ws, r_sec, _C_HEAD, r_sec, _C_END)

        if "OPERATING" in head_name:
            ws.append([])

        sub_total_rows = []; direct_item_rows = []

        # Direct items (e.g. CAPEX items)
        for item in (head.get("items") or []):
            r = ws.max_row + 1
            ws.append(["", "", (item.get("sub_head_of_expense") or "").strip(), item.get("name", ""),
                        *item.get("q1", [0, 0, 0]), *item.get("q2", [0, 0, 0]),
                        *item.get("q3", [0, 0, 0]), *item.get("q4", [0, 0, 0]),
                        *_qtr_fmls(r)])
            _style_annual_row(ws, ws.max_row, font=_NORMAL)
            _fmt_annual_row(ws, ws.max_row)
            direct_item_rows.append(ws.max_row)

        # Sub-heads
        sub_ctr = 1
        for sub in (head.get("sub_heads") or []):
            roman = _to_roman(sub_ctr)
            ws.append(["", roman, sub.get("name", "")])
            r_sub = ws.max_row
            for col in range(_C_HEAD, _C_END + 1):
                c = ws.cell(r_sub, col)
                c.fill = _FILL_ORANGE_LIGHT
                c.font = Font(bold=True, name="Calibri", size=10, color="7A3B00")
                c.border = _BORDER; c.alignment = _LEFT
            ws.cell(r_sub, _C_SI).fill = _FILL_ORANGE_LIGHT
            ws.cell(r_sub, _C_SI).border = _BORDER
            _merge(ws, r_sub, _C_HEAD, r_sub, _C_END)

            sub_item_rows = []
            for item in (sub.get("items") or []):
                r = ws.max_row + 1
                ws.append(["", "", sub.get("name", ""), item.get("name", ""),
                            *item.get("q1", [0, 0, 0]), *item.get("q2", [0, 0, 0]),
                            *item.get("q3", [0, 0, 0]), *item.get("q4", [0, 0, 0]),
                            *_qtr_fmls(r)])
                _style_annual_row(ws, ws.max_row, font=_NORMAL)
                _fmt_annual_row(ws, ws.max_row)
                sub_item_rows.append(ws.max_row)

            if sub_item_rows:
                ws.append(["", "", "", f"TOTAL - {sub.get('name', '')}",
                            *_totals_fmls(sub_item_rows)])
                _style_annual_row(ws, ws.max_row,
                                   PatternFill("solid", fgColor="EBF5FB"), _BOLD)
                _fmt_annual_row(ws, ws.max_row)
                sub_total_rows.append(ws.max_row)
            sub_ctr += 1

        total_rows = sub_total_rows if sub_total_rows else direct_item_rows
        if total_rows:
            ws.append(["", "", "", f"TOTAL - {head.get('name', '')}",
                        *_totals_fmls(total_rows)])
            _style_annual_row(ws, ws.max_row,
                               PatternFill("solid", fgColor="D4E6F1"),
                               Font(bold=True, name="Calibri", size=10, color="003B63"))
            _fmt_annual_row(ws, ws.max_row)
            head_total_rows.append(ws.max_row)
            if "OPERATING" in head_name:
                ws.append([])

    if head_total_rows:
        ws.append(["", "", "", "GRAND TOTAL", *_totals_fmls(head_total_rows)])
        _style_annual_row(ws, ws.max_row, _FILL_BLUE_DARK,
                           Font(bold=True, color="FFFFFF", name="Calibri", size=10))
        _fmt_annual_row(ws, ws.max_row)

    _auto_widths(ws)
    return ws


# ═══════════════════════════════════════════════════════════════════════════════
# SHEET 5 — Estimate Consolidated
#
# Same format as Annual but data comes from group_actuals_ytd… API.
# API shape: items have Q1/Q2/Q3/Q4 keys (float totals) and months dict.
# ═══════════════════════════════════════════════════════════════════════════════

def _sheet_estimate(wb, data, fy, org="Azim Premji Foundation"):
    # JS uses Q1/Q2/Q3/Q4 keys directly (pre-computed by API).
    # months dict is only used when quarters are expanded (individual month cols).
    # We MUST use Q1/Q2/Q3/Q4 as the source of truth to match JS totals exactly.
    MONTH_KEYS = ["4", "5", "6", "7", "8", "9", "10", "11", "12", "1", "2", "3"]
    Q_MONTH_IDX = {"q1": [0,1,2], "q2": [3,4,5], "q3": [6,7,8], "q4": [9,10,11]}

    def _get_quarters(obj):
        """
        Return q1/q2/q3/q4 as 3-element monthly lists.
        Source: months dict for individual values.
        But quarterly TOTAL must equal Q1/Q2/Q3/Q4 keys (JS qTot).
        If months sum != Q key, distribute Q key evenly across 3 months.
        """
        m    = obj.get("months") or {}
        mths = [_fv(m.get(k, 0)) for k in MONTH_KEYS]
        result = {}
        for qk, idxs in Q_MONTH_IDX.items():
            q_key  = qk.upper()   # Q1/Q2/Q3/Q4
            q_val  = _fv(obj.get(q_key, 0))   # authoritative quarter total from API
            m_vals = [mths[i] for i in idxs]
            m_sum  = sum(m_vals)
            if abs(m_sum - q_val) < 0.01:
                # months already add up correctly
                result[qk] = m_vals
            else:
                # months don't match Q key — use Q key split evenly (preserves total)
                # This ensures Excel SUM(monthly) = Q key = matches JS display
                each = q_val / 3.0
                result[qk] = [each, each, q_val - 2 * each]
        return result

    def _norm_obj(obj):
        ql = _get_quarters(obj)
        obj.update(ql)
        for item in (obj.get("items") or []):
            item.update(_get_quarters(item))
        for sub in (obj.get("sub_heads") or []):
            sub.update(_get_quarters(sub))
            for item in (sub.get("items") or []):
                item.update(_get_quarters(item))
        return obj

    normalised = [_norm_obj(dict(h)) for h in (data or [])]

    ws = wb.create_sheet("Estimate Consolidated")
    _sheet_title(ws, org, f"Estimate Consolidated – {fy}", 21)
    _write_annual_headers(ws, fy)

    head_total_rows = []
    head_counter    = 0

    for head in normalised:
        head_counter += 1
        alpha     = chr(64 + head_counter)
        head_name = _norm(head.get("name") or "")

        if "COVID" in head_name:
            ws.append([])
            item = (head.get("items") or [{}])[0]
            r = ws.max_row + 1
            ws.append(["", alpha, head.get("name", ""), item.get("name", ""),
                        *item.get("q1", [0, 0, 0]), *item.get("q2", [0, 0, 0]),
                        *item.get("q3", [0, 0, 0]), *item.get("q4", [0, 0, 0]),
                        *_qtr_fmls(r)])
            _style_annual_row(ws, ws.max_row, font=_NORMAL)
            _fmt_annual_row(ws, ws.max_row)
            # COVID row IS included in GRAND TOTAL (JS sums all heads incl. COVID)
            head_total_rows.append(ws.max_row)
            continue

        ws.append(["", alpha, head.get("name", "")])
        r_sec = ws.max_row
        for col in range(_C_HEAD, _C_END + 1):
            c = ws.cell(r_sec, col)
            c.fill = _FILL_BLUE_LIGHT
            c.font = Font(bold=True, name="Calibri", size=10, color="003B63")
            c.border = _BORDER; c.alignment = _LEFT
        ws.cell(r_sec, _C_SI).fill = _FILL_BLUE_LIGHT
        ws.cell(r_sec, _C_SI).border = _BORDER
        _merge(ws, r_sec, _C_HEAD, r_sec, _C_END)

        if "OPERATING" in head_name:
            ws.append([])

        sub_total_rows = []; direct_item_rows = []

        for item in (head.get("items") or []):
            r = ws.max_row + 1
            ws.append(["", "", (item.get("sub_head_of_expense") or "").strip(), item.get("name", ""),
                        *item.get("q1", [0, 0, 0]), *item.get("q2", [0, 0, 0]),
                        *item.get("q3", [0, 0, 0]), *item.get("q4", [0, 0, 0]),
                        *_qtr_fmls(r)])
            _style_annual_row(ws, ws.max_row, font=_NORMAL)
            _fmt_annual_row(ws, ws.max_row)
            direct_item_rows.append(ws.max_row)

        sub_ctr = 1
        for sub in (head.get("sub_heads") or []):
            roman = _to_roman(sub_ctr)
            ws.append(["", roman, sub.get("name", "")])
            r_sub = ws.max_row
            for col in range(_C_HEAD, _C_END + 1):
                c = ws.cell(r_sub, col)
                c.fill = _FILL_ORANGE_LIGHT
                c.font = Font(bold=True, name="Calibri", size=10, color="7A3B00")
                c.border = _BORDER; c.alignment = _LEFT
            ws.cell(r_sub, _C_SI).fill = _FILL_ORANGE_LIGHT
            ws.cell(r_sub, _C_SI).border = _BORDER
            _merge(ws, r_sub, _C_HEAD, r_sub, _C_END)

            sub_item_rows = []
            for item in (sub.get("items") or []):
                r = ws.max_row + 1
                ws.append(["", "", sub.get("name", ""), item.get("name", ""),
                            *item.get("q1", [0, 0, 0]), *item.get("q2", [0, 0, 0]),
                            *item.get("q3", [0, 0, 0]), *item.get("q4", [0, 0, 0]),
                            *_qtr_fmls(r)])
                _style_annual_row(ws, ws.max_row, font=_NORMAL)
                _fmt_annual_row(ws, ws.max_row)
                sub_item_rows.append(ws.max_row)

            if sub_item_rows:
                ws.append(["", "", "", f"TOTAL - {sub.get('name', '')}",
                            *_totals_fmls(sub_item_rows)])
                _style_annual_row(ws, ws.max_row,
                                   PatternFill("solid", fgColor="EBF5FB"), _BOLD)
                _fmt_annual_row(ws, ws.max_row)
                sub_total_rows.append(ws.max_row)
            sub_ctr += 1

        total_rows = sub_total_rows if sub_total_rows else direct_item_rows
        if total_rows:
            ws.append(["", "", "", f"TOTAL - {head.get('name', '')}",
                        *_totals_fmls(total_rows)])
            _style_annual_row(ws, ws.max_row,
                               PatternFill("solid", fgColor="D4E6F1"),
                               Font(bold=True, name="Calibri", size=10, color="003B63"))
            _fmt_annual_row(ws, ws.max_row)
            head_total_rows.append(ws.max_row)
            if "OPERATING" in head_name:
                ws.append([])

    if head_total_rows:
        ws.append(["", "", "", "GRAND TOTAL", *_totals_fmls(head_total_rows)])
        _style_annual_row(ws, ws.max_row, _FILL_BLUE_DARK,
                           Font(bold=True, color="FFFFFF", name="Calibri", size=10))
        _fmt_annual_row(ws, ws.max_row)

    _auto_widths(ws)
    return ws


# ═══════════════════════════════════════════════════════════════════════════════
# SHEET 6 — Budget & Estimate
#
# JS display: one table, rows = Expense Head / Sub-head / Line Item
#   Columns: Expense Head/Line Item | [per unit: Plan | Est] | Grand Total Plan | Grand Total Est
#   Values in raw rupees (formatINR whole number)
#   Sections from rawData[0].actuals (structure only, values summed across all entities)
#   Total column = GRAND TOTAL section (seq 9999) from each entity
#
# Export format:
#   Row 1 = unit column headers (plan + est per unit + Grand Total plan + Grand Total est)
#   Section rows (blue), sub-head rows (orange), item rows (white)
#   Grand Total row (dark blue)
# ═══════════════════════════════════════════════════════════════════════════════

def _sheet_be(wb, be_data, fy, plan_label, est_label, org="Azim Premji Foundation"):
    ws = wb.create_sheet("Budget & Estimate")
    entities = [e for e in (be_data or []) if not _is_consolidated(e)]

    if not entities:
        ws.cell(1, 1, "No data")
        return ws

    # Build structure from first entity's actuals
    def _build_struct(entity):
        struct = []
        for sec in (entity.get("actuals") or []):
            if _is_gt(sec): continue
            struct.append({
                "name": sec.get("name", ""),
                "sub_heads": [
                    {"name": sub.get("name", ""),
                     "items": [{"name": i.get("name", "")} for i in (sub.get("items") or [])]}
                    for sub in (sec.get("sub_heads") or [])
                ],
                "items": [{"name": i.get("name", "")} for i in (sec.get("items") or [])],
            })
        return struct

    struct = _build_struct(entities[0])

    # Value extraction functions — match JS secVal/subVal/itemVal/grandVal
    def _sec_val(entity, sn, field):
        v = 0.0
        for s in (entity.get("actuals") or []):
            if _is_gt(s) or s.get("name") != sn: continue
            v += _fv(s.get("ytd" if field == "plan" else "total_posted_amt_ytd"))
        return v

    def _sub_val(entity, sn, subn, field):
        v = 0.0
        for s in (entity.get("actuals") or []):
            if _is_gt(s) or s.get("name") != sn: continue
            for sub in (s.get("sub_heads") or []):
                if sub.get("name") != subn: continue
                v += _fv(sub.get("ytd" if field == "plan" else "total_posted_amt_ytd"))
        return v

    def _item_val(entity, nm, field):
        v = 0.0
        fld = "ytd" if field == "plan" else "total_posted_amt"
        for s in (entity.get("actuals") or []):
            if _is_gt(s): continue
            for i in (s.get("items") or []):
                if i.get("name") == nm: v += _fv(i.get(fld))
            for sub in (s.get("sub_heads") or []):
                for i in (sub.get("items") or []):
                    if i.get("name") == nm: v += _fv(i.get(fld))
        return v

    def _grand_val(entity, field):
        # Use GRAND TOTAL section directly (match JS grandVal)
        gt = 0.0; found = False
        fld = "ytd" if field == "plan" else "total_posted_amt_ytd"
        for s in (entity.get("actuals") or []):
            if _is_gt(s):
                gt += _fv(s.get(fld)); found = True
        if not found:
            for s in (entity.get("actuals") or []):
                gt += _fv(s.get(fld))
        return gt

    # Column layout:
    # Col 2 = Expense Head / Line Item
    # Col 3+ = per entity: Plan, Est
    # Last 2 = Grand Total Plan, Grand Total Est
    n_ents   = len(entities)
    c_items  = 2
    c_data   = 3                           # first entity plan col
    c_gt_p   = c_data + n_ents * 2         # Grand Total Plan
    c_gt_e   = c_gt_p + 1                  # Grand Total Est
    last_col = c_gt_e

    _sheet_title(ws, org, f"Budget & Estimate – {fy}", last_col)

    cur = 4

    # ── Header row 1 — entity names ─────────────────────────────────────────
    for col in range(c_items, last_col + 1):
        c = ws.cell(cur, col)
        c.fill = _FILL_BLUE_MID; c.font = _WHITE_BOLD
        c.border = _BORDER; c.alignment = _CENTER

    ws.cell(cur, c_items, "Expense Head / Line Item")
    _merge(ws, cur, c_items, cur + 1, c_items)

    for i, e in enumerate(entities):
        col = c_data + i * 2
        ws.cell(cur, col, (e.get("label") or "").strip())
        _merge(ws, cur, col, cur, col + 1)

    ws.cell(cur, c_gt_p, "Grand Total")
    _merge(ws, cur, c_gt_p, cur, c_gt_e)
    cur += 1

    # ── Header row 2 — Plan / Est per entity ────────────────────────────────
    for col in range(c_items, last_col + 1):
        c = ws.cell(cur, col)
        c.fill = _FILL_ORANGE; c.font = _WHITE_BOLD
        c.border = _BORDER; c.alignment = _CENTER

    for i in range(n_ents):
        col = c_data + i * 2
        ws.cell(cur, col,     plan_label)
        ws.cell(cur, col + 1, est_label)

    ws.cell(cur, c_gt_p, plan_label)
    ws.cell(cur, c_gt_e, est_label)
    cur += 1

    # ── Currency note ────────────────────────────────────────────────────────
    ws.cell(cur, last_col, "Raw ₹")
    ws.cell(cur, last_col).font = _ITALIC
    ws.cell(cur, last_col).alignment = Alignment(horizontal="right")
    cur += 1

    def _fill_data_cols(row, plan_vals, est_vals, gt_p, gt_e,
                        sec_fill=None, sec_font=None):
        for col in range(c_items, last_col + 1):
            c = ws.cell(row, col)
            c.border = _BORDER
            c.alignment = _LEFT if col == c_items else _RIGHT
            if sec_fill: c.fill = sec_fill
            if sec_font: c.font = sec_font
            c.number_format = NUM_FMT_RAW
        for i, (pv, ev) in enumerate(zip(plan_vals, est_vals)):
            col = c_data + i * 2
            ws.cell(row, col,     pv or None).number_format = NUM_FMT_RAW
            ws.cell(row, col + 1, ev or None).number_format = NUM_FMT_RAW
        ws.cell(row, c_gt_p, gt_p or None).number_format = NUM_FMT_RAW
        ws.cell(row, c_gt_e, gt_e or None).number_format = NUM_FMT_RAW

    # ── Data rows ────────────────────────────────────────────────────────────
    for sec in struct:
        sn = sec["name"]

        # Section row — cb-row-head (blue light)
        sec_p_vals = [_sec_val(e, sn, "plan") for e in entities]
        sec_e_vals = [_sec_val(e, sn, "est")  for e in entities]
        gt_sec_p   = sum(sec_p_vals); gt_sec_e = sum(sec_e_vals)

        ws.cell(cur, c_items, sn)
        _fill_data_cols(cur, sec_p_vals, sec_e_vals, gt_sec_p, gt_sec_e,
                        _FILL_BLUE_LIGHT,
                        Font(bold=True, name="Calibri", size=10, color="003B63"))
        cur += 1

        # Sub-head rows — cb-row-sub (orange light)
        for sub in sec.get("sub_heads", []):
            subn = sub["name"]
            sub_p = [_sub_val(e, sn, subn, "plan") for e in entities]
            sub_e = [_sub_val(e, sn, subn, "est")  for e in entities]
            gt_sp = sum(sub_p); gt_se = sum(sub_e)

            ws.cell(cur, c_items, f"  {subn}")
            _fill_data_cols(cur, sub_p, sub_e, gt_sp, gt_se,
                            _FILL_ORANGE_LIGHT,
                            Font(bold=True, name="Calibri", size=10, color="7A3B00"))
            cur += 1

            # Item rows
            for item in sub.get("items", []):
                nm = item["name"]
                ip = [_item_val(e, nm, "plan") for e in entities]
                ie = [_item_val(e, nm, "est")  for e in entities]
                gt_ip = sum(ip); gt_ie = sum(ie)
                ws.cell(cur, c_items, f"    {nm}")
                _fill_data_cols(cur, ip, ie, gt_ip, gt_ie, _FILL_WHITE, _NORMAL)
                cur += 1

        # Direct items on section
        for item in sec.get("items", []):
            nm = item["name"]
            ip = [_item_val(e, nm, "plan") for e in entities]
            ie = [_item_val(e, nm, "est")  for e in entities]
            gt_ip = sum(ip); gt_ie = sum(ie)
            ws.cell(cur, c_items, f"  {nm}")
            _fill_data_cols(cur, ip, ie, gt_ip, gt_ie, _FILL_WHITE, _NORMAL)
            cur += 1

    # Grand Total row — use GRAND TOTAL section (seq 9999) per entity
    gt_p_vals = [_grand_val(e, "plan") for e in entities]
    gt_e_vals = [_grand_val(e, "est")  for e in entities]
    all_gp    = sum(gt_p_vals); all_ge = sum(gt_e_vals)

    ws.cell(cur, c_items, "GRAND TOTAL")
    for col in range(c_items, last_col + 1):
        c = ws.cell(cur, col)
        c.fill = _FILL_BLUE_DARK
        c.font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
        c.border = _BORDER; c.alignment = _LEFT if col == c_items else _RIGHT
        c.number_format = NUM_FMT_RAW
    for i, (pv, ev) in enumerate(zip(gt_p_vals, gt_e_vals)):
        col = c_data + i * 2
        ws.cell(cur, col,     pv or None).number_format = NUM_FMT_RAW
        ws.cell(cur, col + 1, ev or None).number_format = NUM_FMT_RAW
    ws.cell(cur, c_gt_p, all_gp or None).number_format = NUM_FMT_RAW
    ws.cell(cur, c_gt_e, all_ge or None).number_format = NUM_FMT_RAW

    _auto_widths(ws)
    return ws


# ═══════════════════════════════════════════════════════════════════════════════
# WHITELISTED API ENDPOINTS
# ═══════════════════════════════════════════════════════════════════════════════

@frappe.whitelist()
def export_ppt(financial_year, ppt_rows, prev_ppt_rows,
               budget_label, est_label, prev_budget_label, prev_est_label):
    fy = financial_year or "2025-26"
    wb = Workbook(); wb.remove(wb.active)
    _sheet_ppt(wb,
               json.loads(ppt_rows),      budget_label,      est_label,
               json.loads(prev_ppt_rows), prev_budget_label, prev_est_label,
               fy, _prev_fy(fy))
    return {"filename": f"Foundation_Metrics_{fy}.xlsx", "data": _wb_to_b64(wb)}


@frappe.whitelist()
def export_summary_inr(financial_year, summary_data):
    fy = financial_year or "2025-26"
    wb = Workbook(); wb.remove(wb.active)
    _sheet_summary_inr(wb, json.loads(summary_data), fy)
    return {"filename": f"Summary_INR_{fy}.xlsx", "data": _wb_to_b64(wb)}


@frappe.whitelist()
def export_headcount(financial_year, headcount_data):
    fy  = financial_year or "2025-26"
    raw = json.loads(headcount_data)
    wb  = Workbook(); wb.remove(wb.active)
    _sheet_headcount(wb, raw, fy)
    return {"filename": f"Headcount_{fy}.xlsx", "data": _wb_to_b64(wb)}


@frappe.whitelist()
def export_annual(financial_year, annual_data):
    fy = financial_year or "2025-26"
    wb = Workbook(); wb.remove(wb.active)
    _sheet_annual(wb, json.loads(annual_data), fy)
    return {"filename": f"Annual_Budget_Consolidated_{fy}.xlsx", "data": _wb_to_b64(wb)}


@frappe.whitelist()
def export_estimate(financial_year, estimate_data):
    fy = financial_year or "2025-26"
    wb = Workbook(); wb.remove(wb.active)
    _sheet_estimate(wb, json.loads(estimate_data), fy)
    return {"filename": f"Estimate_Consolidated_{fy}.xlsx", "data": _wb_to_b64(wb)}


@frappe.whitelist()
def export_budget_estimate(financial_year, be_data):
    fy     = financial_year or "2025-26"
    labels = _fy_labels(fy)
    wb     = Workbook(); wb.remove(wb.active)
    _sheet_be(wb, json.loads(be_data), fy, labels["plan"], labels["est"])
    return {"filename": f"Budget_and_Estimate_{fy}.xlsx", "data": _wb_to_b64(wb)}


def _sheet_budget_vs_actual(wb, heads, fy, month):
    """Expense Head -> Sub Head -> Item, one row each, Budget/Actual/Util%/
    Variance columns - mirrors the Vue "Budget vs Actual" page's own
    expand-all table exactly (same get_combined_actuals tree, same
    Budget=ytd / Actual=total_posted_amt(_ytd) field reads), so the
    exported file always matches what's on screen when "Export Excel" is
    clicked, indentation and grand total included.
    """
    ws = wb.create_sheet("Budget vs Actual")
    ws.column_dimensions["A"].width = 45
    for col in "BCDE":
        ws.column_dimensions[col].width = 18

    ws["A1"] = f"Budget vs Actual — YTD {month} {fy}"
    ws["A1"].font = Font(bold=True, name="Calibri", size=13)
    ws.merge_cells("A1:E1")

    header_row = 3
    headers = ["Expense Head / Sub Head / Item", "Budget", "Actual", "Util %", "Variance"]
    for col, text in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=text)
        c.font, c.fill, c.alignment, c.border = _WHITE_BOLD, _FILL_BLUE_MID, _CENTER, _BORDER

    def item_actual(item):
        return _fv(item.get("total_posted_amt", item.get("total_posted_amt_ytd", 0)))

    def util(budget, actual):
        return round(actual / budget, 4) if budget else 0

    row = header_row + 1
    grand_budget = grand_actual = 0.0

    for head in heads or []:
        head_budget = _fv(head.get("ytd", 0))
        head_actual = _fv(head.get("total_posted_amt_ytd", 0))
        grand_budget += head_budget
        grand_actual += head_actual

        c = ws.cell(row=row, column=1, value=head.get("name", ""))
        c.font, c.fill, c.border, c.alignment = _BOLD, _FILL_BLUE_LIGHT, _BORDER, _LEFT
        for col, val in ((2, head_budget), (3, head_actual)):
            cell = ws.cell(row=row, column=col, value=val)
            cell.number_format, cell.font, cell.fill, cell.border, cell.alignment = NUM_FMT_RAW, _BOLD, _FILL_BLUE_LIGHT, _BORDER, _RIGHT
        u = ws.cell(row=row, column=4, value=util(head_budget, head_actual))
        u.number_format, u.font, u.fill, u.border, u.alignment = NUM_FMT_PCT, _BOLD, _FILL_BLUE_LIGHT, _BORDER, _RIGHT
        v = ws.cell(row=row, column=5, value=head_budget - head_actual)
        v.number_format, v.font, v.fill, v.border, v.alignment = NUM_FMT_RAW, _BOLD, _FILL_BLUE_LIGHT, _BORDER, _RIGHT
        row += 1

        for sub in head.get("sub_heads", []) or []:
            sub_budget = _fv(sub.get("ytd", 0))
            sub_actual = _fv(sub.get("total_posted_amt_ytd", 0))
            c = ws.cell(row=row, column=1, value="  " + sub.get("name", ""))
            c.fill, c.border, c.alignment = _FILL_ORANGE_LIGHT, _BORDER, _LEFT
            for col, val in ((2, sub_budget), (3, sub_actual)):
                cell = ws.cell(row=row, column=col, value=val)
                cell.number_format, cell.fill, cell.border, cell.alignment = NUM_FMT_RAW, _FILL_ORANGE_LIGHT, _BORDER, _RIGHT
            u = ws.cell(row=row, column=4, value=util(sub_budget, sub_actual))
            u.number_format, u.fill, u.border, u.alignment = NUM_FMT_PCT, _FILL_ORANGE_LIGHT, _BORDER, _RIGHT
            v = ws.cell(row=row, column=5, value=sub_budget - sub_actual)
            v.number_format, v.fill, v.border, v.alignment = NUM_FMT_RAW, _FILL_ORANGE_LIGHT, _BORDER, _RIGHT
            row += 1

            for item in sub.get("items", []) or []:
                item_budget = _fv(item.get("ytd", 0))
                item_act = item_actual(item)
                c = ws.cell(row=row, column=1, value="    " + item.get("name", ""))
                c.border, c.alignment = _BORDER, _LEFT
                for col, val in ((2, item_budget), (3, item_act)):
                    cell = ws.cell(row=row, column=col, value=val)
                    cell.number_format, cell.border, cell.alignment = NUM_FMT_RAW, _BORDER, _RIGHT
                u = ws.cell(row=row, column=4, value=util(item_budget, item_act))
                u.number_format, u.border, u.alignment = NUM_FMT_PCT, _BORDER, _RIGHT
                v = ws.cell(row=row, column=5, value=item_budget - item_act)
                v.number_format, v.border, v.alignment = NUM_FMT_RAW, _BORDER, _RIGHT
                row += 1

        # A head with no sub_heads carries its own items directly (matches
        # the Vue page's tree-walk, which only reads head.items when
        # head.sub_heads is empty).
        if not head.get("sub_heads") and head.get("items"):
            for item in head.get("items", []) or []:
                item_budget = _fv(item.get("ytd", 0))
                item_act = item_actual(item)
                c = ws.cell(row=row, column=1, value="  " + item.get("name", ""))
                c.border, c.alignment = _BORDER, _LEFT
                for col, val in ((2, item_budget), (3, item_act)):
                    cell = ws.cell(row=row, column=col, value=val)
                    cell.number_format, cell.border, cell.alignment = NUM_FMT_RAW, _BORDER, _RIGHT
                u = ws.cell(row=row, column=4, value=util(item_budget, item_act))
                u.number_format, u.border, u.alignment = NUM_FMT_PCT, _BORDER, _RIGHT
                v = ws.cell(row=row, column=5, value=item_budget - item_act)
                v.number_format, v.border, v.alignment = NUM_FMT_RAW, _BORDER, _RIGHT
                row += 1

    c = ws.cell(row=row, column=1, value="GRAND TOTAL")
    c.font, c.fill, c.border, c.alignment = _WHITE_BOLD, _FILL_BLUE_DARK, _BORDER, _LEFT
    for col, val in ((2, grand_budget), (3, grand_actual)):
        cell = ws.cell(row=row, column=col, value=val)
        cell.number_format, cell.font, cell.fill, cell.border, cell.alignment = NUM_FMT_RAW, _WHITE_BOLD, _FILL_BLUE_DARK, _BORDER, _RIGHT
    u = ws.cell(row=row, column=4, value=util(grand_budget, grand_actual))
    u.number_format, u.font, u.fill, u.border, u.alignment = NUM_FMT_PCT, _WHITE_BOLD, _FILL_BLUE_DARK, _BORDER, _RIGHT
    v = ws.cell(row=row, column=5, value=grand_budget - grand_actual)
    v.number_format, v.font, v.fill, v.border, v.alignment = NUM_FMT_RAW, _WHITE_BOLD, _FILL_BLUE_DARK, _BORDER, _RIGHT


@frappe.whitelist()
def export_budget_vs_actual(financial_year, month, heads_data):
    fy = financial_year or "2025-26"
    wb = Workbook(); wb.remove(wb.active)
    _sheet_budget_vs_actual(wb, json.loads(heads_data), fy, month or "March")
    return {"filename": f"Budget_vs_Actual_{fy}_{month or 'March'}.xlsx", "data": _wb_to_b64(wb)}


@frappe.whitelist()
def export_all(financial_year,
               ppt_rows, prev_ppt_rows,
               budget_label, est_label, prev_budget_label, prev_est_label,
               summary_data, headcount_data,
               annual_data, estimate_data, be_data):
    fy     = financial_year or "2025-26"
    labels = _fy_labels(fy)
    wb     = Workbook(); wb.remove(wb.active)

    _sheet_ppt(wb,
               json.loads(ppt_rows),      budget_label,      est_label,
               json.loads(prev_ppt_rows), prev_budget_label, prev_est_label,
               fy, _prev_fy(fy))
    _sheet_summary_inr(wb, json.loads(summary_data),   fy)
    _sheet_headcount(  wb, json.loads(headcount_data),  fy)
    _sheet_annual(     wb, json.loads(annual_data),     fy)
    _sheet_estimate(   wb, json.loads(estimate_data),   fy)
    _sheet_be(         wb, json.loads(be_data),         fy, labels["plan"], labels["est"])

    return {
        "filename": f"Foundation_Consolidated_Budget_{fy}.xlsx",
        "data"    : _wb_to_b64(wb),
    }
