import io
import json
import xml.etree.ElementTree as ET
from collections import defaultdict

import frappe
import requests

from budget_mis.api.adjustments import get_adjustments_month_wise, get_monthly_adjustments
from budget_mis.utils import guest_api, get_allowed_units
from budget_mis.api.response_crypto import is_direct_http_call
from budget_mis.api.erp_utils import get_peoplesoft_prod_credentials, financial_year_label

ACTUALS_CACHE_TTL = 15 * 60


# fetch_actuals_month_wise/fetch_actuals_ytd take no unit param at all and
# are cached+shared across every caller - correct for internal callers
# (which must keep seeing everything and filter it themselves afterward),
# but these are also directly reachable from the live Vue "ERP Actuals"
# page (ErpActuals.vue), so a direct HTTP call previously returned the
# ENTIRE organization's raw ERP actuals to any logged-in user regardless of
# their own Finance user access - the same cross-tenant leak already fixed
# in api_1/grouped_actuals.py, in this older parallel module. Filtering the
# already-computed/cached result on the way OUT (never mutating the cached
# object itself) closes this without touching the shared-cache-then-filter
# behavior internal callers depend on.
def _scope_rows_to_allowed_units(rows, fn):
    if not is_direct_http_call(fn):
        return rows
    allowed_units = get_allowed_units()
    if allowed_units is None:
        return rows
    if not allowed_units:
        return []
    return [r for r in rows if str(r.get("business_unit")) in allowed_units]


# * ==============================================================  PeopleSoft API - Testing server  =====================================================================================
@guest_api
def fetch_actuals_ytd_legacy(fiscal_year, accounting_period):
    try:
        fiscal_year = int(fiscal_year)
        accounting_period = int(accounting_period)
    except (TypeError, ValueError):
        frappe.throw("fiscal_year and accounting_period must be integers")

    month_wise = fetch_actuals_month_wise(fiscal_year, accounting_period)
    if month_wise.get("status") != "success":
        return month_wise

    rows = [r for r in month_wise.get("data", []) if str(r.get("accounting_period")) != "0"]
    rows = _scope_rows_to_allowed_units(rows, fetch_actuals_ytd_legacy)

    return {
        "status": "success",
        "params": {
            "fiscal_year": fiscal_year,
            "accounting_period": accounting_period
        },
        "data": rows
    }


# * ==============================================================  PeopleSoft API - Prod, accounting period without opening balance  =====================================================================================
@guest_api
def fetch_actuals_ytd(fiscal_year, accounting_period):
    try:
        month_wise = fetch_actuals_month_wise(fiscal_year, accounting_period)
    except requests.exceptions.Timeout:
        return {
            "status": "failed",
            "error": "Request timeout while connecting to ERP"
        }
    except Exception:
        frappe.log_error(
            title="PeopleSoft API Error",
            message=frappe.get_traceback()
        )
        return {
            "status": "failed",
            "error": "Unexpected server error"
        }

    if month_wise.get("status") != "success":
        return month_wise

    erp_rows = [r for r in month_wise.get("data", []) if str(r.get("accounting_period")) != "0"]
    frappe_rows = get_monthly_adjustments(financial_year_label(fiscal_year), accounting_period)
    combined_rows = _scope_rows_to_allowed_units(erp_rows + frappe_rows, fetch_actuals_ytd)

    return {
        "status": "success",
        "fiscal_year": fiscal_year,
        "accounting_period": accounting_period,
        "count": len(combined_rows),
        "data": combined_rows
    }


# * ==============================================================  PeopleSoft API - Prod, month-wise (the one live HTTP call)  =====================================================================================
@guest_api
def fetch_actuals_month_wise(fiscal_year, accounting_period):
    cache_key = f"budget_mis:erp_actuals_month_wise:{fiscal_year}:{accounting_period}"
    cached = frappe.cache().get_value(cache_key)
    if cached is not None:
        if is_direct_http_call(fetch_actuals_month_wise):
            return {**cached, "data": _scope_rows_to_allowed_units(cached.get("data", []), fetch_actuals_month_wise)}
        return cached

    result = _fetch_actuals_month_wise_raw(fiscal_year, accounting_period)
    if result.get("status") == "success":
        frappe.cache().set_value(cache_key, result, expires_in_sec=ACTUALS_CACHE_TTL)
    if is_direct_http_call(fetch_actuals_month_wise):
        return {**result, "data": _scope_rows_to_allowed_units(result.get("data", []), fetch_actuals_month_wise)}
    return result


def _fetch_actuals_month_wise_raw(fiscal_year, accounting_period):
    demo_path = frappe.conf.get("demo_actuals_json_path")
    if demo_path:
        return _fetch_actuals_month_wise_demo(demo_path, fiscal_year, accounting_period)

    try:
        username, password = get_peoplesoft_prod_credentials()

        if not username or not password:
            frappe.throw("ERP credentials are not configured")

        base_url = (
            "https://pserp.azimpremjifoundation.org:8053/"
            "PSIGW/RESTListeningConnector/"
            "PSFT_EP/ExecuteQuery.v1/PUBLIC/"
            "Z_MIS_ACTUALS_BY_PERIOD/XMLP/NONFILE"
        )

        api_url = (
            f"{base_url}"
            f"?isconnectedquery=N"
            f"&maxrows=100000"
            f"&prompt_uniquepromptname=FISCAL_YEAR,ACCOUNTING_PERIOD"
            f"&prompt_fieldvalue={fiscal_year},{accounting_period}"
        )

        erp_rows = []

        response = requests.get(
            api_url,
            headers={"Accept": "application/xml"},
            auth=(username, password),
            timeout=120
        )

        if response.status_code != 200:
            frappe.log_error("ERP HTTP Error", response.text[:500])
        else:
            try:
                root = ET.fromstring(response.content)
            except ET.ParseError:
                frappe.log_error("XML Parse Error", response.text[:500])
            else:
                for row_elem in root.iter():
                    if row_elem.tag.lower().endswith("row"):
                        row_data = {
                            child.tag.split("}")[-1].lower(): child.text
                            for child in row_elem
                        }
                        erp_rows.append(row_data)

        frappe_rows = get_adjustments_month_wise(financial_year_label(fiscal_year), accounting_period)

        combined_rows = erp_rows + frappe_rows

        return {
            "status": "success",
            "fiscal_year": fiscal_year,
            "accounting_period": accounting_period,
            "count": len(combined_rows),
            "data": combined_rows
        }

    except requests.exceptions.Timeout:
        frappe.log_error("ERP Timeout", "Request timed out")
        return {
            "status": "failed",
            "error": "Request timeout while connecting to ERP"
        }

    except Exception:
        frappe.log_error("API Error", frappe.get_traceback())
        return {
            "status": "failed",
            "error": "Unexpected server error"
        }


def _fetch_actuals_month_wise_demo(demo_path, fiscal_year, accounting_period):
    try:
        with open(demo_path) as f:
            all_rows = json.load(f)
    except (OSError, ValueError):
        frappe.log_error("Demo Actuals JSON Error", frappe.get_traceback())
        return {
            "status": "failed",
            "error": "Could not read demo actuals JSON file"
        }

    erp_rows = [
        row for row in all_rows
        if str(row.get("fiscal_year")) == str(fiscal_year)
        and int(row.get("accounting_period") or 0) <= int(accounting_period)
    ]

    frappe_rows = get_adjustments_month_wise(financial_year_label(fiscal_year), accounting_period)
    combined_rows = erp_rows + frappe_rows

    return {
        "status": "success",
        "fiscal_year": fiscal_year,
        "accounting_period": accounting_period,
        "count": len(combined_rows),
        "data": combined_rows
    }


# * ==============================================================  PeopleSoft API - grouped by business dimensions  =====================================================================================
@guest_api
def group_actuals_by_dimension(fiscal_year, accounting_period):
    try:
        month_wise = fetch_actuals_month_wise(fiscal_year, accounting_period)
    except requests.exceptions.Timeout:
        frappe.log_error("PeopleSoft Timeout", "ERP request timed out")
        return {"status": "failed", "error": "Request timeout while connecting to ERP"}
    except Exception:
        frappe.log_error("PeopleSoft API Unexpected Error", frappe.get_traceback())
        return {"status": "failed", "error": "Unexpected server error"}

    if month_wise.get("status") != "success":
        return month_wise

    grouped = {}
    for row_data in month_wise.get("data", []):
        if str(row_data.get("accounting_period")) == "0":
            continue

        key = (
            row_data.get("business_unit"),
            row_data.get("account"),
            row_data.get("deptid"),
            row_data.get("operating_unit"),
            row_data.get("accounting_period")
        )

        amt = float(row_data.get("posted_total_amt") or 0)

        if key not in grouped:
            grouped[key] = {
                "business_unit": key[0],
                "account": key[1],
                "deptid": key[2],
                "operating_unit": key[3],
                "accounting_period": key[4],
                "posted_total_amt": amt
            }
        else:
            grouped[key]["posted_total_amt"] += amt

    result = list(grouped.values())

    return {
        "status": "success",
        "fiscal_year": fiscal_year,
        "accounting_period": accounting_period,
        "count": len(result),
        "data": result
    }


# * ==============================================================  PeopleSoft API - GL-code level breakdown  =====================================================================================
@guest_api
def group_actuals_by_gl_code(fiscal_year, accounting_period):
    month_wise = group_actuals_month_wise(fiscal_year, accounting_period)

    if month_wise.get("status") != "success":
        return {"status": month_wise.get("status", "failed"), "fiscal_year": fiscal_year, "data": []}

    grouped = defaultdict(float)

    for record in month_wise.get("data", []):
        key = (
            record.get("business_unit"),
            record.get("deptid"),
            record.get("operating_unit"),
            record.get("head_of_expense"),
            record.get("sub_head_of_expense"),
            record.get("type_of_expense"),
            record.get("sequence_id"),
            record.get("account"),
        )
        grouped[key] += float(record.get("total_posted_amt") or 0)

    final_output = []
    for (
        business_unit,
        deptid,
        operating_unit,
        head_of_expense,
        sub_head_of_expense,
        type_of_expense,
        sequence_id,
        account,
    ), total_sum in grouped.items():
        final_output.append({
            "sequence_id": sequence_id,
            "business_unit": business_unit,
            "deptid": deptid,
            "operating_unit": operating_unit,
            "head_of_expense": head_of_expense,
            "sub_head_of_expense": sub_head_of_expense,
            "type_of_expense": type_of_expense,
            "gl_code": account,
            "total_posted_amt": round(total_sum, 2)
        })

    final_output.sort(key=lambda x: (x.get("sequence_id") or 0, x.get("gl_code") or ""))

    return {
        "status": "success",
        "fiscal_year": fiscal_year,
        "data": final_output
    }


# * ==============================================================  PeopleSoft API - grouped actuals (YTD, GL codes collapsed)  =====================================================================================
@guest_api
def group_actuals_ytd(fiscal_year, accounting_period):
    month_wise = group_actuals_month_wise(fiscal_year, accounting_period)

    if month_wise.get("status") != "success":
        return {"status": month_wise.get("status", "failed"), "fiscal_year": fiscal_year, "data": []}

    grouped = defaultdict(float)
    meta_by_key = {}

    for record in month_wise.get("data", []):
        key = (
            record.get("business_unit"),
            record.get("deptid"),
            record.get("operating_unit"),
            record.get("head_of_expense"),
            record.get("sub_head_of_expense"),
            record.get("type_of_expense"),
            record.get("sequence_id"),
        )
        grouped[key] += float(record.get("total_posted_amt") or 0)
        meta_by_key.setdefault(key, record.get("actuals_type_of_expenses"))

    final_output = []
    for (
        business_unit,
        deptid,
        operating_unit,
        head_of_expense,
        sub_head_of_expense,
        type_of_expense,
        sequence_id,
    ), total_sum in grouped.items():
        final_output.append({
            "sequence_id": sequence_id,
            "business_unit": business_unit,
            "deptid": deptid,
            "operating_unit": operating_unit,
            "head_of_expense": head_of_expense,
            "sub_head_of_expense": sub_head_of_expense,
            "type_of_expense": type_of_expense,
            "actuals_type_of_expenses": meta_by_key.get((
                business_unit, deptid, operating_unit, head_of_expense,
                sub_head_of_expense, type_of_expense, sequence_id,
            )) or type_of_expense,
            "total_posted_amt": round(total_sum, 2)
        })

    final_output.sort(key=lambda x: (x.get("sequence_id") or 0))

    return {
        "status": "success",
        "fiscal_year": fiscal_year,
        "data": final_output
    }


# * ==============================================================  PeopleSoft API - grouped actuals, month-wise (Expenses / GL code Mapping join)  =====================================================================================
@guest_api
def group_actuals_month_wise(fiscal_year, accounting_period):
    erp_response = group_actuals_by_dimension(fiscal_year, accounting_period)

    if "message" in erp_response:
        erp_data = erp_response.get("message", {}).get("data", [])
    else:
        erp_data = erp_response.get("data", [])

    if not erp_data:
        return {
            "status": "success",
            "fiscal_year": fiscal_year,
            "accounting_period": accounting_period,
            "data": [],
            "count": 0
        }

    expenses = frappe.get_all(
        "Expenses",
        fields=[
            "name",
            "head_of_expense",
            "sub_head_of_expense",
            "type_of_expense",
            "sequence_id"
        ]
    )

    expense_lookup = {str(e.name): e for e in expenses}

    child_rows = frappe.get_all(
        "GL code Mapping",
        fields=["parent", "gl_code_map"]
    )

    gl_parent_map = {str(row.gl_code_map).strip(): str(row.parent) for row in child_rows}

    grouped = defaultdict(float)

    for record in erp_data:
        account = str(record.get("account", "")).strip()
        period = record.get("accounting_period")
        amount = float(record.get("posted_total_amt", 0) or 0)

        parent = gl_parent_map.get(account)
        if not parent:
            continue

        expense = expense_lookup.get(parent)
        if not expense:
            continue

        key = (
            period,
            record.get("business_unit"),
            record.get("deptid"),
            record.get("operating_unit"),
            account,
            expense.head_of_expense,
            expense.sub_head_of_expense,
            expense.type_of_expense,
            expense.sequence_id
        )

        grouped[key] += amount

    final_output = []

    for (
        period,
        business_unit,
        deptid,
        operating_unit,
        account,
        head_of_expense,
        sub_head_of_expense,
        type_of_expense,
        sequence_id
    ), total_sum in grouped.items():

        final_output.append({
            "sequence_id": sequence_id,
            "accounting_period": period,
            "account": account,
            "business_unit": business_unit,
            "deptid": deptid,
            "operating_unit": operating_unit,
            "head_of_expense": head_of_expense,
            "sub_head_of_expense": sub_head_of_expense,
            "type_of_expense": type_of_expense,
            "actuals_type_of_expenses": type_of_expense,
            "total_posted_amt": round(total_sum, 2)
        })

    final_output.sort(key=lambda x: (x.get("sequence_id") or 0))

    return {
        "status": "success",
        "fiscal_year": fiscal_year,
        "accounting_period": accounting_period,
        "data": final_output,
        "count": len(final_output)
    }


# * ==============================================================  Excel export for the ERP Actuals page  =====================================================================================
ERP_ACTUALS_EXPORT_METHODS = {
    "fetch_actuals_ytd": fetch_actuals_ytd,
    "fetch_actuals_month_wise": fetch_actuals_month_wise,
}


@frappe.whitelist()
def export_actuals_excel(fiscal_year, accounting_period, rows=None, method_key=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    if rows is not None:
        rows = frappe.parse_json(rows) if isinstance(rows, str) else rows
    else:
        fn = ERP_ACTUALS_EXPORT_METHODS.get(method_key)
        if not fn:
            frappe.throw("Invalid export method")

        result = fn(fiscal_year=fiscal_year, accounting_period=accounting_period)

        if result.get("status") != "success":
            frappe.throw(result.get("error") or "ERP request failed")

        rows = result.get("data") or []
        # fn() is an INTERNAL call here (export_actuals_excel is the one
        # actually dispatched over HTTP, under its own cmd name) - fn's own
        # is_direct_http_call check never fires for this call path, so the
        # export would otherwise bypass the scoping just added to
        # fetch_actuals_ytd/fetch_actuals_month_wise and download every
        # unit's actuals regardless of the caller's own Finance user access.
        rows = _scope_rows_to_allowed_units(rows, export_actuals_excel)

    # fiscal_year is dropped — same as the on-screen table, it just repeats
    # the fiscal_year/accounting_period already shown for the whole export.
    columns = []
    for row in rows:
        for key in (row or {}).keys():
            if key != "fiscal_year" and key not in columns:
                columns.append(key)

    wb = Workbook()
    ws = wb.active
    ws.title = "ERP Actuals"

    header_fill = PatternFill("solid", fgColor="0076B6")
    header_font = Font(bold=True, color="FFFFFF")

    ws.append(columns)
    for col_idx in range(1, len(columns) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font

    # Amount columns get real Excel numbers (right-aligned, thousands separator);
    # code-like columns (account, deptid, accounting_period, ...) stay as text so
    # leading zeros / exact codes aren't mangled by Excel's numeric coercion.
    AMOUNT_COLUMNS = {"posted_total_amt"}
    AMOUNT_FORMAT = "#,##0.00;[Red]-#,##0.00"

    for row in rows:
        values = []
        for col in columns:
            value = row.get(col)
            if col == "is_adjustment":
                value = "True" if str(value).strip().lower() in ("1", "true", "yes") else "False"
            elif col in AMOUNT_COLUMNS:
                value = frappe.utils.flt(value)
            values.append(value)
        ws.append(values)
        if AMOUNT_COLUMNS:
            for col_idx, col in enumerate(columns, start=1):
                if col in AMOUNT_COLUMNS:
                    ws.cell(row=ws.max_row, column=col_idx).number_format = AMOUNT_FORMAT

    if "posted_total_amt" in columns:
        total = sum(frappe.utils.flt(row.get("posted_total_amt")) for row in rows)
        total_row = ["" for _ in columns]
        total_row[0] = "Total"
        total_row[columns.index("posted_total_amt")] = total
        ws.append(total_row)

        total_fill = PatternFill("solid", fgColor="EAF3FA")
        total_font = Font(bold=True, color="0076B6")
        for col_idx, col in enumerate(columns, start=1):
            cell = ws.cell(row=ws.max_row, column=col_idx)
            cell.fill = total_fill
            cell.font = total_font
            if col in AMOUNT_COLUMNS:
                cell.number_format = AMOUNT_FORMAT

    for col_idx, col in enumerate(columns, start=1):
        max_len = max([len(col)] + [len(str(r.get(col, "") or "")) for r in rows]) if rows else len(col)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)

    stream = io.BytesIO()
    wb.save(stream)

    frappe.response["filename"] = f"ERP_Actuals_{fiscal_year}_P{accounting_period}.xlsx"
    frappe.response["filecontent"] = stream.getvalue()
    frappe.response["type"] = "binary"


# * ==============================================================  Old names kept for the Desk page (main/page/erp_actuals)  =====================================================================================
get_actuals_from_erp_prod = fetch_actuals_ytd
get_actuals_from_erp_month_wise = fetch_actuals_month_wise
export_erp_actuals_excel = export_actuals_excel
